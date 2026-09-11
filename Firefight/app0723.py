import os
import json
import difflib
import re
from typing import Any, Dict, List, Optional, Tuple

import asyncio
import httpx
import requests
import boto3
from dotenv import load_dotenv
from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.responses import JSONResponse, FileResponse
from starlette.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from starlette.staticfiles import StaticFiles

# ============== 环境 ==============
load_dotenv()

VERBOSE_LOG = os.getenv("VERBOSE_LOG", "0") == "1"

def debug_log(*args, **kwargs):
    if VERBOSE_LOG:
        print(*args, **kwargs)

TOS_ENDPOINT = os.getenv("TOS_ENDPOINT", "tos-cn-beijing.volces.com")
TOS_BUCKET = os.getenv("TOS_BUCKET", "nanjing-fire")
TOS_REGION = os.getenv("TOS_REGION", "cn-beijing")
TOS_ACCESS_KEY = os.getenv("TOS_ACCESS_KEY", "")
TOS_SECRET_KEY = os.getenv("TOS_SECRET_KEY", "")

DOUBAO_API_URL = os.getenv("DOUBAO_API_URL", "https://ark.cn-beijing.volces.com/api/v3/responses")
DOUBAO_API_KEY = os.getenv("DOUBAO_API_KEY", "")
DOUBAO_MODEL = os.getenv("DOUBAO_MODEL", "doubao-seed-2-0-pro-260215")

# StageA 的一级、二级、三级候选白名单来自问题分级.xlsx。
ISSUE_CATALOG_XLSX = os.getenv("ISSUE_CATALOG_XLSX", "问题分级.xlsx")
ISSUE_CATALOG_SHEET = os.getenv("ISSUE_CATALOG_SHEET", "sheet1")
ISSUE_CATALOG_HEADER_ROW = int(os.getenv("ISSUE_CATALOG_HEADER_ROW", "1"))
ISSUE_COL_CAT1 = os.getenv("ISSUE_COL_CAT1", "一级")
ISSUE_COL_CAT2 = os.getenv("ISSUE_COL_CAT2", "二级")
ISSUE_COL_CAT3 = os.getenv("ISSUE_COL_CAT3", "三级（具体问题）")
ISSUE_COL_REGULATION = os.getenv("ISSUE_COL_REGULATION", "条例")

KB_DOMAIN = os.getenv("KB_DOMAIN", "api-knowledgebase.mlp.cn-beijing.volces.com")
KB_SERVICE_RESOURCE_ID = os.getenv("KB_SERVICE_RESOURCE_ID", "kb-service-58c1523ebd033284")
KB_API_KEY = os.getenv("KB_API_KEY", "")

KB_HTTP_TIMEOUT = int(os.getenv("KB_HTTP_TIMEOUT", "30"))
KB_PER_TYPE_TOPK = int(os.getenv("KB_PER_TYPE_TOPK", "1"))

# 独立规范条文知识库：案例库只负责返回条文号，规范库按 record_id 精确返回原文。
CLAUSE_KB_ENABLED = int(os.getenv("CLAUSE_KB_ENABLED", "1"))
CLAUSE_KB_DOMAIN = os.getenv("CLAUSE_KB_DOMAIN", KB_DOMAIN)
CLAUSE_KB_SERVICE_RESOURCE_ID = os.getenv("CLAUSE_KB_SERVICE_RESOURCE_ID", "")
CLAUSE_KB_API_KEY = os.getenv("CLAUSE_KB_API_KEY", KB_API_KEY)
CLAUSE_KB_CONCURRENCY = max(1, int(os.getenv("CLAUSE_KB_CONCURRENCY", "1")))
CLAUSE_KB_SCAN_LIMIT = max(1, int(os.getenv("CLAUSE_KB_SCAN_LIMIT", "15")))
CLAUSE_KB_MAX_EXACT_MATCHES = max(1, int(os.getenv("CLAUSE_KB_MAX_EXACT_MATCHES", "1")))
CLAUSE_KB_CACHE_ENABLED = int(os.getenv("CLAUSE_KB_CACHE_ENABLED", "1"))

# —— 并发与超时 —— #
FS_STAGEA_ROUNDS = int(os.getenv("FS_STAGEA_ROUNDS", "3"))
FS_FUSION_ROUNDS = int(os.getenv("FS_FUSION_ROUNDS", "3"))
FS_MAX_TYPES     = max(5, int(os.getenv("FS_MAX_TYPES", "5")))
FS_TEMP_LIST     = [float(x.strip()) for x in os.getenv("FS_TEMP_LIST", "0.6,0.8,1.0").split(",") if x.strip()]
FS_MODEL_CONCURRENCY = int(os.getenv("FS_MODEL_CONCURRENCY", "4"))
FS_KB_CONCURRENCY    = int(os.getenv("FS_KB_CONCURRENCY", "5"))
FS_HTTP_TIMEOUT      = int(os.getenv("FS_HTTP_TIMEOUT", "300"))

# —— Stage0 开关与阈值（新增） —— #
FS_STAGE0_ENABLED   = int(os.getenv("FS_STAGE0_ENABLED", "1"))
FS_STAGE0_THRESHOLD = float(os.getenv("FS_STAGE0_THRESHOLD", "0.6"))

if not all([TOS_BUCKET, TOS_ACCESS_KEY, TOS_SECRET_KEY, DOUBAO_API_KEY]):
    raise RuntimeError("缺少环境变量：TOS_* / DOUBAO_API_KEY / TOS_BUCKET")

# ============== FastAPI ==============
app = FastAPI(title="FireSafety Async Ensemble + Stage0")
app.mount("/static", StaticFiles(directory="static"), name="static")

@app.get("/")
def _index():
    return FileResponse("static/index.html")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"],
)

# ============== TOS ==============
tos_client = boto3.client(
    "s3",
    endpoint_url=f"https://{TOS_ENDPOINT}",
    aws_access_key_id=TOS_ACCESS_KEY,
    aws_secret_access_key=TOS_SECRET_KEY,
    region_name=TOS_REGION,
)

def _tos_fixed_host() -> str:
    return f"{TOS_BUCKET}.{TOS_ENDPOINT}"

# 用预签名 PUT URL 上传图片
def upload_image_to_tos(file: UploadFile) -> str:
    ext = os.path.splitext(file.filename or "")[-1].lower()
    if ext not in [".png", ".jpg", ".jpeg", ".bmp", ".webp"]:
        ext = ".jpg" if not ext else ext
    key = f"safe_uploads_stageA/{os.urandom(8).hex()}{ext}"
    try:
        content_type = file.content_type or "application/octet-stream"
        put_url = tos_client.generate_presigned_url(
            ClientMethod="put_object",
            Params={"Bucket": TOS_BUCKET, "Key": key, "ContentType": content_type},
            ExpiresIn=600
        )
        from urllib.parse import urlparse
        parsed = urlparse(put_url)
        put_url = put_url.replace(parsed.netloc, _tos_fixed_host())

        resp = requests.put(put_url, data=file.file.read(), headers={"Content-Type": content_type}, timeout=30)
        if resp.status_code != 200:
            raise HTTPException(status_code=500, detail=f"TOS 上传失败：{resp.status_code}")
        return key
    except HTTPException:
        raise
    except Exception as e:
        print("[ERROR] TOS 上传异常:", e)
        raise HTTPException(status_code=500, detail="TOS 上传失败")


# 为对象生成 GET 临时 URL
def _presigned_get_url(key: str, expires: int = 600) -> str:
    try:
        url = tos_client.generate_presigned_url(
            ClientMethod="get_object",
            Params={"Bucket": TOS_BUCKET, "Key": key},
            ExpiresIn=expires
        )
        from urllib.parse import urlparse
        parsed = urlparse(url)
        return url.replace(parsed.netloc, _tos_fixed_host())
    except Exception as e:
        print("[ERROR] 生成 GET 签名失败:", e)
        raise HTTPException(status_code=500, detail="生成签名URL失败")

# ============== Excel 三级分类白名单 ==============
_ISSUE_ID_RE = re.compile(r"(?P<id>\d+(?:\.\d+)+-\d+)")
_CODE_NAME_RE = re.compile(r"^\s*(?P<code>\d+(?:\.\d+)*)[.、]?\s*(?P<name>.*)$")
_STANDARD_CODE_RE = re.compile(
    r"(?<![A-Z0-9])(?P<code>GB(?:\s*/\s*T)?\s*\d{4,6}(?:\.\d+)?\s*[-—–－−]\s*\d{4})(?![A-Z0-9])",
    re.IGNORECASE,
)
_CLAUSE_NO_RE = re.compile(r"(?<![\d.])(?P<clause>\d+(?:\.\d+){1,3})(?![\d.])")
_CLAUSE_KB_CACHE: Dict[str, List[dict]] = {}
_KB_CATEGORY_MAPPING = None
_KB_ALLOWED_TEXT: str = None


def _cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _split_code_name(value: str) -> Tuple[str, str]:
    text = _cell_text(value)
    match = _CODE_NAME_RE.match(text)
    if not match:
        return "", text
    return match.group("code"), match.group("name").strip(" .、")


def _load_category_mapping_from_excel(xlsx_path: str, sheet: str) -> dict:
    """读取问题分级表，形成 issue_id -> 一级/二级/三级固定映射。"""
    try:
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
        if sheet not in wb.sheetnames:
            raise ValueError(f"找不到工作表：{sheet}")
        ws = wb[sheet]

        headers = {}
        for col in range(1, ws.max_column + 1):
            value = _cell_text(ws.cell(row=ISSUE_CATALOG_HEADER_ROW, column=col).value)
            if value:
                headers[value] = col

        required = [ISSUE_COL_CAT1, ISSUE_COL_CAT2, ISSUE_COL_CAT3]
        missing = [name for name in required if name not in headers]
        if missing:
            raise ValueError(f"缺少列：{', '.join(missing)}")

        c1_idx = headers[ISSUE_COL_CAT1]
        c2_idx = headers[ISSUE_COL_CAT2]
        c3_idx = headers[ISSUE_COL_CAT3]
        regulation_idx = headers.get(ISSUE_COL_REGULATION)

        issues_by_id = {}
        hierarchy = {}
        last_raw1 = ""
        last_raw2 = ""

        for row in range(ISSUE_CATALOG_HEADER_ROW + 1, ws.max_row + 1):
            current_raw1 = _cell_text(ws.cell(row=row, column=c1_idx).value)
            current_raw2 = _cell_text(ws.cell(row=row, column=c2_idx).value)
            raw3 = _cell_text(ws.cell(row=row, column=c3_idx).value)
            regulation = (
                _cell_text(ws.cell(row=row, column=regulation_idx).value)
                if regulation_idx else ""
            )

            # 兼容一级、二级分类使用合并单元格的情况。
            if current_raw1:
                if current_raw1 != last_raw1:
                    last_raw2 = ""
                last_raw1 = current_raw1
            if current_raw2:
                last_raw2 = current_raw2

            raw1, raw2 = last_raw1, last_raw2
            if not raw3:
                continue

            match = _ISSUE_ID_RE.search(raw3)
            if not match:
                debug_log(f"[WARN] 第{row}行无法解析三级问题编号：{raw3}")
                continue

            issue_id = match.group("id")
            code1, name1 = _split_code_name(raw1)
            code2, name2 = _split_code_name(raw2)

            name3 = raw3[:match.start()] + raw3[match.end():]
            name3 = re.sub(
                r"^\s*问题\s*Q?\s*[-—:：]?\s*",
                "",
                name3,
                flags=re.IGNORECASE,
            )
            name3 = name3.strip(" .、-—：:")
            if not name3:
                continue

            item = {
                "category1_code": code1,
                "category1": name1 or raw1,
                "category2_code": code2,
                "category2": name2 or raw2,
                "issue_id": issue_id,
                "category3": name3,
                "regulation_index_raw": regulation,
            }
            if issue_id in issues_by_id:
                debug_log(f"[WARN] 重复 issue_id，保留首次记录：{issue_id}")
                continue

            issues_by_id[issue_id] = item
            hierarchy.setdefault(code1 or item["category1"], {
                "category1_code": code1,
                "category1": item["category1"],
                "category2": {},
            })
            level2 = hierarchy[code1 or item["category1"]]["category2"].setdefault(
                code2 or item["category2"],
                {
                    "category2_code": code2,
                    "category2": item["category2"],
                    "issues": [],
                },
            )
            level2["issues"].append(item)

        if not issues_by_id:
            raise ValueError("问题分级表中没有可用的三级问题")

        for level1 in hierarchy.values():
            for level2 in level1["category2"].values():
                level2["issues"].sort(key=lambda x: x["issue_id"])

        return {
            "issues_by_id": issues_by_id,
            "hierarchy": hierarchy,
        }
    except HTTPException:
        raise
    except Exception as e:
        print("[ERROR] 加载三级分类白名单失败：", e)
        raise HTTPException(status_code=500, detail=f"服务器未正确加载问题分级表：{e}")


def _render_allowed_categories_text(catalog: dict) -> str:
    lines = ["一级 / 二级 / 三级问题白名单："]
    for level1 in catalog["hierarchy"].values():
        c1 = " ".join(x for x in [level1["category1_code"], level1["category1"]] if x)
        lines.append(f"一级：{c1}")
        for level2 in level1["category2"].values():
            c2 = " ".join(x for x in [level2["category2_code"], level2["category2"]] if x)
            lines.append(f"  二级：{c2}")
            for issue in level2["issues"]:
                lines.append(f"    - {issue['issue_id']}｜{issue['category3']}")
    return "\n".join(lines)


# 保留旧函数名，避免 main.py 或其他模块调用处需要同步修改。
def _get_kb_categories() -> Tuple[dict, str]:
    global _KB_CATEGORY_MAPPING, _KB_ALLOWED_TEXT
    if _KB_CATEGORY_MAPPING is None:
        if not os.path.exists(ISSUE_CATALOG_XLSX):
            raise HTTPException(status_code=500,detail=f"ISSUE_CATALOG_XLSX 不存在：{ISSUE_CATALOG_XLSX}",)
        _KB_CATEGORY_MAPPING = _load_category_mapping_from_excel(ISSUE_CATALOG_XLSX,ISSUE_CATALOG_SHEET,)
        _KB_ALLOWED_TEXT = _render_allowed_categories_text(_KB_CATEGORY_MAPPING)
        debug_log("[INFO] 已加载三级问题白名单：",len(_KB_CATEGORY_MAPPING["issues_by_id"]),"个三级问题",)
    return _KB_CATEGORY_MAPPING, _KB_ALLOWED_TEXT


# ============== 通用工具 ==============
def _extract_json(text: str) -> str:
    import re
    m = re.search(r"```json\s*(\{.*\})\s*```", text, re.DOTALL)
    if m:
        return m.group(1).strip()
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end > start:
        return text[start:end + 1].strip()
    return text.strip()

# 构造多模态输入，兼容两种常见返回文本位置
def _build_responses_payload(prompt_text: str, image_urls: List[str] = None, temperature: float = None) -> dict:
    image_urls = image_urls or []
    content = []
    for image_url in image_urls:
        if image_url:
            content.append({"type": "input_image", "image_url": image_url})
    content.append({"type": "input_text", "text": prompt_text})
    payload = {
        "model": DOUBAO_MODEL,
        "input": [
            {
                "role": "user",
                "content": content
            }
        ]
    }
    if temperature is not None:
        payload["temperature"] = float(temperature)
    return payload

def _extract_response_text(data: dict) -> str:
    if not data:
        return ""
    if isinstance(data.get("output_text"), str):
        return data["output_text"]
    texts = []
    for item in data.get("output", []) or []:
        for c in item.get("content", []) or []:
            if c.get("type") == "output_text" and c.get("text"):
                texts.append(c["text"])
            elif c.get("type") == "text" and c.get("text"):
                texts.append(c["text"])
    return "\n".join(texts).strip()

def _closest(subcand: str, options: list, cutoff: float = 0.65):
    if not subcand:
        return None
    match = difflib.get_close_matches(subcand.strip(), options or [], n=1, cutoff=cutoff)
    return match[0] if match else None

def _clean_kb_text(t: str) -> str:
    if not t:
        return ""
    return t.replace("<KBImage>", "").replace("</KBImage>", "").strip()

def _normalize_ws(s: str) -> str:
    return " ".join((s or "").split())

def _violations_match_kb(hazard: dict, kb_text: str) -> bool:
    """只校验知识库中已经提供的规范原文，不校验条款索引。"""
    kb_norm = _normalize_ws(kb_text)
    for v in hazard.get("violations", []) or []:
        content = _normalize_ws(v.get("content", ""))
        if not content or content not in kb_norm:
            return False
    return True

def _compact_for_match(value: str) -> str:
    """用于比对规范名称、条款号，忽略空格和中英文标点差异。"""
    return "".join(ch.lower() for ch in (value or "") if ch.isalnum())

def _regulation_ref_is_in_kb(ref: dict, kb_text: str) -> bool:
    """判断条款索引是否来自本次检索结果；不要求知识库已有条文原文。"""
    standard = _compact_for_match(str(ref.get("standard", "")))
    clause = _compact_for_match(str(ref.get("clause", "")))
    kb_compact = _compact_for_match(_clean_kb_text(kb_text))
    return bool(standard and clause and standard in kb_compact and clause in kb_compact)

# 条款索引去重、标记核验状态；并检查规范号和条款号是否在本次 KB 文本内
def _prepare_regulation_refs(hazard: dict, kb_text: str) -> List[dict]:
    """标准化条款索引，并始终标记为待专业核对。"""
    refs = []
    seen = set()
    for raw in hazard.get("regulation_refs", []) or []:
        if not isinstance(raw, dict):
            continue
        standard = str(raw.get("standard", "")).strip()
        clause = str(raw.get("clause", "")).strip()
        if not standard or not clause:
            continue
        key = (_compact_for_match(standard), _compact_for_match(clause))
        if key in seen:
            continue
        seen.add(key)
        refs.append({
            "standard": standard,
            "clause": clause,
            "verification_status": "待专业核对",
            "source": "知识库条款索引" if _regulation_ref_is_in_kb(raw, kb_text) else "模型推断",
        })
    return refs

# 仅保留条文内容能在 KB 返回文本中逐字出现的记录，避免模型生成原文被当成已核验规范
def _verified_violations_from_kb(hazard: dict, kb_text: str) -> List[dict]:
    """仅保留能在本次知识库文本中逐字找到的规范原文。"""
    verified = []
    kb_norm = _normalize_ws(kb_text)
    for raw in hazard.get("violations", []) or []:
        if not isinstance(raw, dict):
            continue
        content = _normalize_ws(str(raw.get("content", "")))
        if not content or content not in kb_norm:
            continue
        verified.append({
            "standard": str(raw.get("standard", "")).strip(),
            "clause": str(raw.get("clause", "")).strip(),
            "content": str(raw.get("content", "")).strip(),
            "verification_status": "已核对原文",
        })
    return verified

# ============== 【新增】Stage0 场景筛查（异步） ==============
def _build_stage0_prompt() -> str:
    # 低温度、严格 JSON。给出正例元素，尽量抑制“万物皆有隐患”的错判。
    return (
        "请判断这张图片是否与【建筑消防验收、消防巡检、建筑施工质量检查或消防设施检查】场景相关。\n"
        "【判断原则】\n"
        "1. 本阶段只判断场景相关性，不判断图片中是否存在消防隐患。\n"
        "2. 不判断图片质量、清晰度、拍摄角度、遮挡情况，也不判断是否能够仅凭图片确认具体问题。\n"
        "3. 图片中不一定会直接出现消防设备。只要属于建筑内部、建筑外部、施工现场、"
        "机电安装区域或消防设施可能设置的位置，均应判为相关。\n"
        "4. 无法确定是否相关时，应优先判为相关，交由后续阶段继续分析，不要轻易拦截。\n\n"
        "【应判为不相关的场景】\n"
        "仅当图片明显属于以下内容且与建筑、施工或消防检查没有关系时，才判为不相关：\n"
        "- 自然风景、人物自拍、食物、动物；\n"
        "- 纯商品图片、纯商标、纯图标、与消防无关的文字截图；\n"
        "- 普通车辆内部、生活用品或其他明显非建筑场景。\n\n"
        " 【相关场景通常包含但不限于】\n"
        "- 建筑室内、走廊、房间、楼梯间、前室、地下室、设备间、管井、吊顶内部；\n"
        "- 建筑外部或周边环境/走廊/房间/楼梯/门窗/墙面/天花、管线/风管/电缆桥架、喷淋/消火栓/消防水带/灭火器、泵房/配电柜/控制柜、应急照明/疏散指示;\n"
        "- 建筑外立面、屋面、消防车道、消防登高操作场地、建筑周边；\n"
        "- 墙体、门窗、楼板、吊顶、洞口、防火分隔和防火封堵位置；\n"
        "- 消防管道、给排水管道、喷淋、消火栓、消防水带、灭火器；\n"
        "- 风管、防烟排烟设施、电缆桥架、线缆、配电柜、控制柜；\n"
        "- 火灾报警设备、应急照明、疏散指示、防火门、消防电梯；\n"
        "- 尚未完工、设备缺失、设施损坏、局部构件或近距离拍摄的建筑施工图片等。\n\n"
        "仅输出严格 JSON（不要代码块、不要解释）：\n"
        "{\n"
        '  "relevant": true,\n'
        '  "confidence": 0.0,\n'
        '  "scene_tags": ["室内","走廊","喷淋","配电柜","管线"],\n'
        '  "reason": "一句话说明判断依据"\n'
        "}\n"
    )

def _parse_stage0_json(raw: str) -> dict:
    try:
        data = json.loads(_extract_json(raw))
    except Exception:
        return {"relevant": True, "confidence": 1.0, "scene_tags": [], "reason": "parse_fallback_true"}
    rel = bool(data.get("relevant", False))
    conf = float(data.get("confidence", 0.0) or 0.0)
    tags = data.get("scene_tags", []) or []
    reason = str(data.get("reason", "") or "")
    # 纠偏：若模型输出 relevant 但信心很低，保持原值；若输出不相关且有明显建筑标签，也不强行改。
    return {"relevant": rel, "confidence": max(0.0, min(conf, 1.0)), "scene_tags": tags, "reason": reason}

# 用豆包判断图片是否与建筑消防巡检有关；返回 relevant、confidence、scene_tags、reason。API/解析失败时默认“放行”。
async def run_stage0_guard_async(client: httpx.AsyncClient, sem: asyncio.Semaphore, image_url: str) -> dict:
    payload = _build_responses_payload(
        prompt_text=_build_stage0_prompt(),
        image_urls=[image_url],
        temperature=0.2
    )
    headers = {"Content-Type": "application/json", "Authorization": f"Bearer {DOUBAO_API_KEY}"}
    async with sem:
        r = await client.post(DOUBAO_API_URL, headers=headers, json=payload)
    if r.status_code != 200:
        # 失败时不阻断流程（维持旧行为）：默认放行
        print(f"[Stage0] 调用失败：HTTP {r.status_code} {r.text[:120]}")
        return {"relevant": True, "confidence": 1.0, "scene_tags": [], "reason": "api_fallback_true"}
    raw = _extract_response_text(r.json())
    out = _parse_stage0_json(raw)
    debug_log(f"[Stage0] 判定：relevant={out['relevant']} conf={out['confidence']:.2f} tags={out['scene_tags'][:4]} reason={out['reason']}")
    return out


# ============== StageA 无白名单命中时的开放式兜底 ==============
def _build_catalog_unmatched_prompt() -> str:
    """StageA没有任何有效白名单候选时，开放式判断图片是否确有疑似问题。"""
    return (
        "你是一名消防验收与消防巡检图像分析助手。\n"
        "前一阶段已确认该图片属于建筑、施工或消防检查相关场景，"
        "但系统未能在既有问题分级表中匹配到任何有效三级问题。\n\n"
        "请暂时不要受问题分级表限制，直接根据图片中能够观察到的事实进行开放式复核：\n"
        "1. 判断图片中是否存在明确或较高可能性的消防、建筑防火、消防设施安装或施工质量问题；\n"
        "2. 只能依据图片可见事实，不得把正常构造、正常设备或无法确认的内容强行判为隐患；\n"
        "3. 若没有足够证据，应将 suspected_hazard 设为 false；\n"
        "4. 若存在疑似问题，可给出1～3条模型建议项，但不要编造问题编号；\n"
        "5. 不要填写规范名称、条款号或规范原文，规范依据后续必须由专业人员核验；\n"
        "6. 建议项应说明可见位置、现场现象、判断依据和建议人工检查内容。\n\n"
        "仅输出严格JSON，不要解释、不要代码块：\n"
        "{\n"
        '  "suspected_hazard": true,\n'
        '  "confidence": 0.0,\n'
        '  "summary": "对图片的总体判断",\n'
        '  "suggestions": [\n'
        "    {\n"
        '      "suggested_issue": "模型建议的问题名称，不填写issue_id",\n'
        '      "point": "图片中的具体位置或对象",\n'
        '      "description": "结合可见事实说明疑似问题",\n'
        '      "visual_evidence": ["可见依据1", "可见依据2"],\n'
        '      "recommended_action": "建议人工复核或补拍的内容"\n'
        "    }\n"
        "  ],\n"
        '  "reason": "为什么判断为有问题或无问题"\n'
        "}\n"
    )


def _parse_catalog_unmatched_result(raw_text: str) -> dict:
    """解析开放式兜底结果；解析失败时不误判为无隐患，而是标记需人工复核。"""
    try:
        data = json.loads(_extract_json(raw_text))
    except Exception as exc:
        return {
            "analysis_available": False,
            "suspected_hazard": False,
            "confidence": 0.0,
            "summary": "",
            "suggestions": [],
            "reason": f"fallback_parse_error: {exc}",
        }

    suspected = bool(data.get("suspected_hazard", False))
    try:
        confidence = float(data.get("confidence", 0.0) or 0.0)
    except (TypeError, ValueError):
        confidence = 0.0
    confidence = max(0.0, min(confidence, 1.0))

    suggestions = []
    if suspected:
        for raw in data.get("suggestions", []) or []:
            if not isinstance(raw, dict):
                continue
            suggested_issue = str(raw.get("suggested_issue", "") or "").strip()
            point = str(raw.get("point", "") or "").strip()
            description = str(raw.get("description", "") or "").strip()
            evidence = [
                str(item).strip()
                for item in (raw.get("visual_evidence", []) or [])
                if str(item).strip()
            ][:5]
            recommended_action = str(raw.get("recommended_action", "") or "").strip()

            # 至少要有问题名称、位置或描述之一，避免保存空建议项。
            if not any([suggested_issue, point, description]):
                continue
            suggestions.append({
                "suggested_issue": suggested_issue or "未分类疑似问题",
                "point": point,
                "description": description,
                "visual_evidence": evidence,
                "recommended_action": recommended_action,
                "catalog_match_status": "unmatched",
                "review_required": True,
            })
            if len(suggestions) >= 3:
                break

    # 模型声称存在问题却没有给出可用建议时，视为结果不完整，需要人工复核。
    analysis_available = not (suspected and not suggestions)
    return {
        "analysis_available": analysis_available,
        "suspected_hazard": suspected,
        "confidence": confidence,
        "summary": str(data.get("summary", "") or "").strip(),
        "suggestions": suggestions,
        "reason": str(data.get("reason", "") or "").strip(),
    }


async def run_catalog_unmatched_fallback_async(
    client: httpx.AsyncClient,
    sem: asyncio.Semaphore,
    image_url: str,
) -> dict:
    """当StageA没有白名单候选时，调用模型做一次不受白名单限制的开放式复核。"""
    try:
        raw = await _doubao_chat_image(
            client,
            sem,
            image_url,
            _build_catalog_unmatched_prompt(),
            temperature=0.2,
        )
        result = _parse_catalog_unmatched_result(raw)
    except Exception as exc:
        print(f"[StageA-Fallback] 开放式复核失败：{type(exc).__name__}: {exc}")
        result = {
            "analysis_available": False,
            "suspected_hazard": False,
            "confidence": 0.0,
            "summary": "",
            "suggestions": [],
            "reason": f"fallback_api_error: {exc}",
        }

    debug_log(
        "[StageA-Fallback] "
        f"available={result.get('analysis_available')} "
        f"suspected={result.get('suspected_hazard')} "
        f"suggestions={len(result.get('suggestions', []))}"
    )
    return result


def build_catalog_unmatched_final_result(fallback_result: dict) -> dict:
    """将开放式兜底结果转换为前端和MongoDB统一使用的final_result。"""
    available = bool(fallback_result.get("analysis_available", False))
    suspected = bool(fallback_result.get("suspected_hazard", False))
    suggestions = fallback_result.get("suggestions", []) or []

    if not available:
        return {
            "analysis_status": "catalog_unmatched",
            "catalog_match_status": "unmatched",
            "hazard_count": 0,
            "hazards": [],
            "unmatched_suggestions": [],
            "review_required": True,
            "review_reason": (
                "图片已通过消防巡检场景筛查，但StageA未在问题分级表中匹配到有效问题；"
                "开放式模型复核未能得到完整结果，请人工复核。"
            ),
            "analysis_message": "未在问题分级表里匹配到对应问题。",
            "fallback_summary": fallback_result.get("summary", ""),
            "fallback_reason": fallback_result.get("reason", ""),
            "fallback_confidence": fallback_result.get("confidence", 0.0),
        }

    if suspected:
        return {
            "analysis_status": "catalog_unmatched",
            "catalog_match_status": "unmatched",
            "hazard_count": 0,
            "hazards": [],
            "unmatched_suggestions": suggestions,
            "review_required": True,
            "review_reason": (
                "图片通过消防巡检场景筛查，模型开放式复核认为存在疑似问题，"
                "但未在问题分级表里匹配到对应问题，请人工复核并考虑补充问题分级表。"
            ),
            "analysis_message": "未在问题分级表里匹配到对应问题，以下为模型建议项。",
            "fallback_summary": fallback_result.get("summary", ""),
            "fallback_reason": fallback_result.get("reason", ""),
            "fallback_confidence": fallback_result.get("confidence", 0.0),
        }

    return {
        "analysis_status": "no_hazard",
        "catalog_match_status": "no_stageA_match",
        "hazard_count": 0,
        "hazards": [],
        "unmatched_suggestions": [],
        "review_required": False,
        "review_reason": "",
        "analysis_message": (
            "图片已通过消防巡检场景筛查；StageA未匹配到白名单问题，"
            "开放式复核也未发现有充分图像证据支持的消防隐患。"
        ),
        "fallback_summary": fallback_result.get("summary", ""),
        "fallback_reason": fallback_result.get("reason", ""),
        "fallback_confidence": fallback_result.get("confidence", 0.0),
    }


# ============== StageA（异步）：直接输出一级/二级/三级候选 ==============
def _build_stageA_prompt(allowed_text: str, run_tag: str = "") -> str:
    tag = f"\n【抽样序号】{run_tag}\n" if run_tag else ""
    return (
        "你是一名消防隐患图像分析助手。请基于现场图片，识别【至少5条，最多6条】疑似或者确定的隐患。\n"
        "每条隐患必须从下面的一级/二级/三级白名单中选择，必须填写白名单中的 issue_id，严禁自造编号。\n"
        "同一个 issue_id 不得重复。优先选择图片中可见的隐患；为了保持高召回，即使只有1类明显隐患，也要结合相近视觉特征、同一设施常见问题或伴生问题补足到至少5条候选。\n\n"
        f"{allowed_text}\n\n"
        "仅输出 JSON，不要解释或代码块：\n"
        "{\n"
        '  "hazard_queries": [\n'
        "    {\n"
        '      "issue_id": "从白名单选择，例如1.1.2-7",\n'
        '      "category1": "对应一级分类",\n'
        '      "category2": "对应二级分类",\n'
        '      "category3": "对应三级问题",\n'
        '      "keywords": ["1~4个现场视觉关键词"]\n'
        "    }\n"
        "  ]\n"
        "}\n\n"
        "若图中只明显出现1类隐患，请在同一大类下给出4~5个【最可能的小类】（基于相近特征或常见伴生问题），确保总数≥5。"
        + tag
    )


async def _doubao_chat_image(
    client: httpx.AsyncClient,
    sem: asyncio.Semaphore,
    image_url: str,
    prompt_text: str,
    temperature: float = 0.7,
) -> str:
    payload = _build_responses_payload(
        prompt_text=prompt_text,
        image_urls=[image_url],
        temperature=temperature,
    )
    headers = {"Content-Type": "application/json", "Authorization": f"Bearer {DOUBAO_API_KEY}"}
    async with sem:
        r = await client.post(DOUBAO_API_URL, headers=headers, json=payload)
    if r.status_code != 200:
        raise HTTPException(status_code=500, detail=f"StageA 调用失败：{r.text[:200]}")
    return _extract_response_text(r.json())


def _parse_stageA_types(raw_text: str, catalog: dict) -> List[Dict[str, str]]:
    """只接受白名单中的 issue_id；一级、二级、三级名称全部从 Excel 主数据回填。"""
    raw = _extract_json(raw_text)
    data = json.loads(raw)
    hqs = data.get("hazard_queries", []) or []
    valid, seen = [], set()

    for item in hqs:
        raw_issue_id = str(item.get("issue_id", "")).strip()
        match = _ISSUE_ID_RE.search(raw_issue_id)
        issue_id = match.group("id") if match else raw_issue_id

        master = catalog["issues_by_id"].get(issue_id)
        if not master or issue_id in seen:
            continue

        candidate = dict(master)
        candidate["keywords"] = [
            str(x).strip()
            for x in (item.get("keywords") or [])
            if str(x).strip()
        ][:4]
        valid.append(candidate)
        seen.add(issue_id)

        if len(valid) >= 8:
            break

    return valid


# 三轮并发采样，按 issue_id 出现频次排序，最终保留前 FS_MAX_TYPES（默认5）个三级候选。
async def ensemble_stageA_async(image_url: str,allowed_text: str,mapping: dict,rounds: int,
                                temps: List[float],model_sem: asyncio.Semaphore,client: httpx.AsyncClient,):
    debug_log(f"[Ensemble-StageA] 并发采样：rounds={rounds}, temps={temps}")
    tasks = []
    temperatures = temps or [0.7]
    for i in range(rounds):
        temp = temperatures[i % len(temperatures)]
        tag = f"R{i + 1}-T{temp}"
        prompt = _build_stageA_prompt(allowed_text, run_tag=tag)
        tasks.append(_doubao_chat_image(client, model_sem, image_url, prompt, temperature=temp))

    texts = await asyncio.gather(*tasks, return_exceptions=True)
    freq: Dict[str, int] = {}
    candidate_by_id: Dict[str, dict] = {}
    per_round: List[List[Dict[str, str]]] = []

    for idx, result in enumerate(texts, 1):
        if isinstance(result, Exception):
            print(f"[StageA] 第{idx}轮失败：{result}")
            per_round.append([])
            continue
        try:
            types_i = _parse_stageA_types(result, mapping)
        except Exception as e:
            print(f"[StageA] 第{idx}轮解析失败：{e}")
            per_round.append([])
            continue

        per_round.append(types_i)
        debug_log(f"[StageA] 第{idx}轮识别：")
        for candidate in types_i:
            issue_id = candidate["issue_id"]
            freq[issue_id] = freq.get(issue_id, 0) + 1
            candidate_by_id.setdefault(issue_id, candidate)
            debug_log(
                f"  - {candidate['category1']} - {candidate['category2']} - "
                f"{issue_id} {candidate['category3']}"
            )

    ranked_ids = sorted(freq, key=lambda issue_id: (-freq[issue_id], issue_id))
    votes = []
    for issue_id in ranked_ids:
        vote = dict(candidate_by_id[issue_id])
        vote["count"] = freq[issue_id]
        votes.append(vote)

    debug_log("[StageA] 三级候选频次统计：")
    for vote in votes:
        debug_log(f"  {vote['issue_id']} {vote['category3']}：{vote['count']}/{rounds}")

    selected = [dict(candidate_by_id[issue_id]) for issue_id in ranked_ids[:FS_MAX_TYPES]]
    return selected, votes, per_round


# ============== KB（异步） ==============
def _kb_headers():
    if not all([KB_SERVICE_RESOURCE_ID, KB_API_KEY]):
        raise HTTPException(status_code=500, detail="知识库配置缺失（KB_SERVICE_RESOURCE_ID/KB_API_KEY）")
    return {
        "Accept": "application/json",
        "Content-Type": "application/json;charset=UTF-8",
        "Host": KB_DOMAIN,
        "Authorization": f"Bearer {KB_API_KEY}",
    }


def _build_kb_query(candidate: dict) -> str:
    path = " > ".join(
        x for x in [candidate.get("category1", ""), candidate.get("category2", ""), candidate.get("category3", "")]
        if x
    )
    return (
        f"问题编号：{candidate.get('issue_id', '')}\n"
        f"隐患路径：{path}\n"
        f"视觉关键词：{'、'.join(candidate.get('keywords', []) or [])}\n"
        "请返回与该三级问题高度相关的：\n"
        "- 隐患描述（专业、简洁）；\n"
        "- 对应规范的名称和条款号；\n"
        "- 如知识库已收录，再返回可核对的规范原文；\n"
        "- 示例图片（如有）。\n"
        "当前知识库可能只有条款索引；条款索引不是规范原文，不能补写或臆造原文。"
    )


# 向配置的远程知识库 chat 服务请求检索，不读本地向量库。
async def kb_chat_async(client: httpx.AsyncClient, sem: asyncio.Semaphore, query_text: str) -> dict:
    url = f"http://{KB_DOMAIN}/api/knowledge/service/chat"

    payload = {
        "service_resource_id": KB_SERVICE_RESOURCE_ID,
        "messages": [{"role": "user", "content": query_text}],
        "stream": False,
    }
    headers = _kb_headers()
    async with sem:
        r = await client.post(url, headers=headers, json=payload)
    if r.status_code != 200:
        raise HTTPException(status_code=500, detail=f"KB 检索失败：HTTP {r.status_code}")
    data = r.json()
    if data.get("code") != 0:
        raise HTTPException(status_code=500, detail=f"KB 业务错误：{data}")
    return data


# 对每个三级候选并发检索；抽取远端 content、score、图片附件链接。
async def search_kb_for_types_async(
    types: List[Dict[str, str]],
    topk: int,
    kb_sem: asyncio.Semaphore,
    client: httpx.AsyncClient,
):
    tasks = [kb_chat_async(client, kb_sem, _build_kb_query(candidate)) for candidate in types]
    packs = await asyncio.gather(*tasks, return_exceptions=True)
    results = []

    identity_fields = [
        "category1_code", "category1", "category2_code", "category2",
        "issue_id", "category3", "keywords", "regulation_index_raw",
    ]

    for candidate, pack_result in zip(types, packs):
        base = {key: candidate.get(key, "") for key in identity_fields}
        if isinstance(pack_result, Exception):
            print(f"[KB] {candidate['issue_id']} 失败：{pack_result}")
            results.append({**base, "results": []})
            continue

        result_list = (((pack_result or {}).get("data") or {}).get("result_list")) or []
        evidence = []
        for entry in result_list[:topk]:
            text = entry.get("content", "") or ""
            score = float(entry.get("score", 0.0) or 0.0)
            images = []
            for attachment in entry.get("chunk_attachment", []) or []:
                if attachment.get("type") == "image" and attachment.get("link"):
                    images.append(attachment["link"])
            evidence.append({"content": text, "images": images, "score": score})

        results.append({**base, "results": evidence})

    return results


# ============== 规范条文知识库（按 record_id 精确关联） ==============
def _normalize_standard_code(value: Any) -> str:
    """GB 50016-2014 / GB/T 12345-2020 -> GB50016-2014 / GB/T12345-2020。"""
    text_value = str(value or "").upper()
    text_value = (
        text_value.replace("—", "-")
        .replace("–", "-")
        .replace("－", "-")
        .replace("−", "-")
    )
    match = _STANDARD_CODE_RE.search(text_value)
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group("code"))


def _normalize_clause_no(value: Any) -> str:
    """第6.2.9条 -> 6.2.9。"""
    match = _CLAUSE_NO_RE.search(str(value or ""))
    return match.group("clause") if match else ""


def _normalize_clause_lookup(value: Any) -> Tuple[str, str]:
    """将四级编号按规范库主条文方式关联：6.2.9.3 -> 6.2.9，子项3。"""
    requested = _normalize_clause_no(value)
    if not requested:
        return "", ""
    parts = requested.split(".")
    if len(parts) >= 4:
        return ".".join(parts[:3]), ".".join(parts[3:])
    return requested, ""


def _make_clause_ref(standard: Any, clause: Any, source: str) -> Optional[dict]:
    standard_code = _normalize_standard_code(standard)
    requested_clause_no = _normalize_clause_no(clause)
    clause_no, subitem_no = _normalize_clause_lookup(requested_clause_no)
    if not standard_code or not clause_no:
        return None
    return {
        "standard_code": standard_code,
        "clause_no": clause_no,
        "requested_clause_no": requested_clause_no,
        "subitem_no": subitem_no,
        "record_id": f"{standard_code}_{clause_no}",
        "source": source,
    }


def _parse_refs_from_text(value: Any, source: str) -> List[dict]:
    """从问题分级表或案例库条款索引文本中提取规范号和条文号。"""
    text_value = str(value or "")
    code_matches = list(_STANDARD_CODE_RE.finditer(text_value))
    refs: List[dict] = []
    seen = set()

    for index, code_match in enumerate(code_matches):
        segment_end = code_matches[index + 1].start() if index + 1 < len(code_matches) else len(text_value)
        segment = text_value[code_match.end():segment_end]
        for clause_match in _CLAUSE_NO_RE.finditer(segment):
            ref = _make_clause_ref(code_match.group("code"), clause_match.group("clause"), source)
            if not ref or ref["record_id"] in seen:
                continue
            seen.add(ref["record_id"])
            refs.append(ref)
    return refs


def _collect_hazard_regulation_refs(hazard: dict) -> List[dict]:
    """优先使用案例融合提取出的条文号，同时用问题分级表固定依据兜底。"""
    refs: List[dict] = []
    seen = set()

    for raw in hazard.get("regulation_refs", []) or []:
        if not isinstance(raw, dict):
            continue
        standard = raw.get("standard_code") or raw.get("standard") or raw.get("standard_title")
        clause = raw.get("requested_clause_no") or raw.get("clause_no") or raw.get("clause")
        ref = _make_clause_ref(standard, clause, raw.get("source") or "案例知识库条款索引")
        if not ref or ref["record_id"] in seen:
            continue
        seen.add(ref["record_id"])
        refs.append(ref)

    for ref in _parse_refs_from_text(
        hazard.get("regulation_index_raw", ""),
        source="问题分级表候选规范依据",
    ):
        if ref["record_id"] in seen:
            continue
        seen.add(ref["record_id"])
        refs.append(ref)

    return refs


def _knowledge_service_url(domain: str) -> str:
    clean_domain = str(domain or "").strip().rstrip("/")
    if clean_domain.startswith("http://") or clean_domain.startswith("https://"):
        return f"{clean_domain}/api/knowledge/service/chat"
    return f"https://{clean_domain}/api/knowledge/service/chat"


def _clause_kb_headers() -> dict:
    missing = []
    if not CLAUSE_KB_DOMAIN:
        missing.append("CLAUSE_KB_DOMAIN")
    if not CLAUSE_KB_SERVICE_RESOURCE_ID:
        missing.append("CLAUSE_KB_SERVICE_RESOURCE_ID")
    if not CLAUSE_KB_API_KEY:
        missing.append("CLAUSE_KB_API_KEY")
    if missing:
        raise RuntimeError(f"规范知识库缺少配置：{', '.join(missing)}")
    return {
        "Accept": "application/json",
        "Content-Type": "application/json;charset=UTF-8",
        "Authorization": f"Bearer {CLAUSE_KB_API_KEY}",
    }


async def clause_kb_chat_async(
    client: httpx.AsyncClient,
    sem: asyncio.Semaphore,
    record_id: str,
) -> dict:
    """使用 record_id 查询规范知识库，并对429自动退避重试。"""
    payload = {
        "service_resource_id": CLAUSE_KB_SERVICE_RESOURCE_ID,
        "messages": [
            {
                "role": "user",
                "content": record_id,
            }
        ],
        "stream": False,
    }

    url = _knowledge_service_url(CLAUSE_KB_DOMAIN)
    headers = _clause_kb_headers()
    response = None

    for attempt in range(5):
        async with sem:
            response = await client.post(
                url,
                headers=headers,
                json=payload,
            )

        if response.status_code != 429:
            break
        retry_after = response.headers.get( "Retry-After", "",).strip()

        try:
            wait_seconds = float(retry_after)
        except (TypeError, ValueError):
            wait_seconds = min(1.5 * (2 ** attempt), 12,)
        wait_seconds = max(wait_seconds, 1.5)
        print( f"[ClauseKB] {record_id} 触发限流，" f"{wait_seconds:.1f}秒后重试 " f"({attempt + 1}/5)")
        await asyncio.sleep(wait_seconds)

    if response is None:
        raise RuntimeError("规范知识库未返回响应")
    if response.status_code != 200:
        raise RuntimeError(f"规范知识库调用失败：HTTP " f"{response.status_code}")
    data = response.json()
    if data.get("code") != 0:
        raise RuntimeError( "规范知识库业务错误：" f"{data.get('message', '')}" )
    return data


def _kb_result_entries(pack: dict) -> List[dict]:
    data = (pack or {}).get("data") or {}
    entries = data.get("result_list") if isinstance(data, dict) else None
    return [entry for entry in (entries or []) if isinstance(entry, dict)]


def _entry_searchable_text(entry: dict) -> str:
    parts = []
    for key in ("content", "fields", "metadata", "document_fields"):
        value = entry.get(key)
        if value is None:
            continue
        if isinstance(value, str):
            parts.append(value)
        else:
            try:
                parts.append(json.dumps(value, ensure_ascii=False))
            except Exception:
                parts.append(str(value))
    return "\n".join(parts)


def _extract_entry_identity(entry: dict) -> dict:
    text_value = _entry_searchable_text(entry)

    record_match = re.search(
        r"(?im)^\s*record_id\s*[:：]\s*([^\s,，]+)",
        text_value,
    )
    standard_match = re.search(
        r"(?im)^\s*standard_code\s*[:：]\s*([^\r\n,，]+)",
        text_value,
    )
    clause_match = re.search(
        r"(?im)^\s*(?:clause_no|条文编号)\s*[:：]\s*([^\r\n,，]+)",
        text_value,
    )

    record_id = re.sub(r"\s+", "", record_match.group(1)).upper() if record_match else ""
    standard_code = _normalize_standard_code(standard_match.group(1)) if standard_match else ""
    clause_no = _normalize_clause_no(clause_match.group(1)) if clause_match else ""

    if not standard_code:
        standard_code = _normalize_standard_code(text_value)
    if record_id and "_" in record_id:
        record_standard, record_clause = record_id.rsplit("_", 1)
        standard_code = standard_code or record_standard
        clause_no = clause_no or _normalize_clause_no(record_clause)

    return {
        "record_id": record_id,
        "standard_code": standard_code,
        "clause_no": clause_no,
    }


def _entry_exactly_matches_ref(entry: dict, ref: dict) -> bool:
    identity = _extract_entry_identity(entry)
    target_record_id = ref["record_id"].upper()
    if identity["record_id"]:
        return identity["record_id"] == target_record_id
    return (
        identity["standard_code"] == ref["standard_code"]
        and identity["clause_no"] == ref["clause_no"]
    )


_CLAUSE_METADATA_LINE_RE = re.compile(
    r"(?im)^\s*(?:record_id|standard_code|standard_title|clause_no|source_url)\s*[:：]"
)


def _strip_clause_metadata_suffix(value: str) -> str:
    text_value = str(value or "").strip()
    match = _CLAUSE_METADATA_LINE_RE.search(text_value)
    return text_value[:match.start()].strip() if match else text_value


def _entry_nested_field(entry: dict, field_name: str) -> str:
    direct = entry.get(field_name)
    if direct not in (None, ""):
        return str(direct).strip()
    for container_name in ("fields", "metadata", "document_fields"):
        container = entry.get(container_name)
        if isinstance(container, dict) and container.get(field_name) not in (None, ""):
            return str(container.get(field_name)).strip()
    return ""


def _parse_clause_entry(entry: dict, ref: dict, retrieval_rank: int) -> dict:
    content = str(entry.get("content") or _entry_nested_field(entry, "content") or "")
    content = content.replace("\r\n", "\n").replace("\r", "\n").strip()

    title_match = re.search(
        r"(?im)^\s*(?:content\s*[：:]\s*)?规范名称\s*[：:]\s*(.+)$",
        content,
    )
    clause_match = re.search(r"(?im)^\s*条文编号\s*[：:]\s*(.+)$", content)
    body_match = re.search(r"(?im)^\s*正文\s*[：:]\s*", content)

    clause_text = ""
    explanation_text = ""
    if body_match:
        body = content[body_match.end():]
        explanation_match = re.search(r"(?im)^\s*条文说明\s*[：:]\s*", body)
        if explanation_match:
            clause_text = _strip_clause_metadata_suffix(body[:explanation_match.start()])
            explanation_text = _strip_clause_metadata_suffix(body[explanation_match.end():])
        else:
            clause_text = _strip_clause_metadata_suffix(body)

    standard_title = (
        title_match.group(1).strip()
        if title_match
        else _entry_nested_field(entry, "standard_title")
    )
    standard_code = (
        _normalize_standard_code(_entry_nested_field(entry, "standard_code"))
        or _normalize_standard_code(standard_title)
        or ref["standard_code"]
    )
    clause_no = (
        _normalize_clause_no(_entry_nested_field(entry, "clause_no"))
        or _normalize_clause_no(clause_match.group(1) if clause_match else "")
        or ref["clause_no"]
    )

    try:
        retrieval_score = float(entry.get("score", 0) or 0)
    except (TypeError, ValueError):
        retrieval_score = 0.0

    return {
        "standard_title": standard_title,
        "standard_code": standard_code,
        "clause_no": clause_no,
        "requested_clause_no": ref.get("requested_clause_no", clause_no),
        "subitem_no": ref.get("subitem_no", ""),
        "clause_text": clause_text,
        "explanation_text": explanation_text,
        "record_id": ref["record_id"],
        "source_url": _entry_nested_field(entry, "source_url"),
        "retrieval_score": retrieval_score,
        "retrieval_rank": retrieval_rank,
        "verification_status": "kb_exact_match",
    }


async def _query_one_clause_ref(
    ref: dict,
    clause_sem: asyncio.Semaphore,
    client: httpx.AsyncClient,
) -> dict:
    record_id = ref["record_id"]
    if CLAUSE_KB_CACHE_ENABLED and record_id in _CLAUSE_KB_CACHE:
        return {
            "unavailable": False,
            "verified": [dict(item) for item in _CLAUSE_KB_CACHE[record_id]],
        }

    try:
        response = await clause_kb_chat_async(client, clause_sem, record_id)
    except Exception as exc:
        print(f"[ClauseKB] {record_id} 调用失败：{type(exc).__name__}: {exc}")
        return {"unavailable": True, "verified": []}

    entries = _kb_result_entries(response)[:CLAUSE_KB_SCAN_LIMIT]
    verified = []
    for rank, entry in enumerate(entries, 1):
        exact = _entry_exactly_matches_ref(entry, ref)
        identity = _extract_entry_identity(entry)
        debug_log(
            f"[ClauseKB] target={record_id} rank={rank} "
            f"returned={identity.get('record_id') or identity.get('standard_code')} "
            f"clause={identity.get('clause_no')} exact={exact}"
        )
        if not exact:
            continue
        parsed = _parse_clause_entry(entry, ref, rank)
        if not parsed["clause_text"]:
            debug_log(f"[ClauseKB] {record_id} 精确命中但正文为空，跳过")
            continue
        verified.append(parsed)
        if len(verified) >= CLAUSE_KB_MAX_EXACT_MATCHES:
            break

    if verified and CLAUSE_KB_CACHE_ENABLED:
        _CLAUSE_KB_CACHE[record_id] = [dict(item) for item in verified]
    return {"unavailable": False, "verified": verified}


async def enrich_hazards_with_clause_kb_async(
    hazards: List[dict],
    clause_sem: asyncio.Semaphore,
    client: httpx.AsyncClient,
) -> Tuple[List[dict], List[dict]]:
    """为最终隐患补充规范原文；只查询最终保留结果，不查询全部StageA候选。"""
    refs_by_issue = {
        hazard.get("issue_id", ""): _collect_hazard_regulation_refs(hazard)
        for hazard in hazards
    }

    if not CLAUSE_KB_ENABLED:
        summaries = []
        for hazard in hazards:
            issue_id = hazard.get("issue_id", "")
            refs = refs_by_issue.get(issue_id, [])
            hazard["regulation_refs"] = refs
            hazard["clause_kb_status"] = "disabled" if refs else "missing_reference"
            hazard["regulation_review_status"] = "规范知识库已关闭" if refs else "未提取到规范条款索引"
            summaries.append({
                "issue_id": issue_id,
                "status": hazard["clause_kb_status"],
                "regulation_refs": refs,
                "verified_count": 0,
            })
        return hazards, summaries

    unique_refs: Dict[str, dict] = {}
    for refs in refs_by_issue.values():
        for ref in refs:
            unique_refs.setdefault(ref["record_id"], ref)

    results_by_record_id: Dict[str, dict] = {}
    # 默认单并发顺序查询，降低规范服务QPS限流风险。
    for index, (record_id, ref) in enumerate(  unique_refs.items()):
        if index > 0:
            await asyncio.sleep(1.0)
        results_by_record_id[record_id] = (
            await _query_one_clause_ref( ref, clause_sem, client, ))

    summaries = []
    for hazard in hazards:
        issue_id = hazard.get("issue_id", "")
        refs = refs_by_issue.get(issue_id, [])
        verified = []
        matched_ref_count = 0
        unavailable = False
        seen = set()

        for ref in refs:
            result = results_by_record_id.get(ref["record_id"], {"unavailable": True, "verified": []})
            unavailable = unavailable or bool(result.get("unavailable"))
            items = result.get("verified", [])
            if items:
                matched_ref_count += 1
            for item in items:
                key = item.get("record_id") or ref["record_id"]
                if key in seen:
                    continue
                seen.add(key)
                verified.append(item)

        if not refs:
            status = "missing_reference"
            review_status = "案例库和问题分级表均未提供可解析的规范号、条文号"
        elif unavailable:
            status = "clause_kb_unavailable"
            review_status = "规范知识库调用失败，条文原文尚未核验"
        elif matched_ref_count == len(refs):
            status = "kb_exact_match"
            review_status = "规范原文已精确匹配，具体适用性仍需专业核验"
        elif matched_ref_count:
            status = "partial_match"
            review_status = "部分规范原文已匹配，其余条文未找到"
        else:
            status = "not_found"
            review_status = "已获得条文索引，但规范知识库未返回精确原文"

        hazard["regulation_refs"] = refs
        # 兼容旧前端字段：violations只放经过规范知识库精确核验的原文。
        hazard["violations"] = [
            {
                "standard": item.get("standard_title") or item.get("standard_code", ""),
                "clause": item.get("requested_clause_no") or item.get("clause_no", ""),
                "content": item.get("clause_text", ""),
                "verification_status": "已核对原文",
            }
            for item in verified
            if item.get("clause_text")
        ]
        hazard["clause_kb_status"] = status
        hazard["regulation_review_status"] = review_status
        hazard["regulation_text_verified"] = bool(verified)

        summaries.append({
            "issue_id": issue_id,
            "status": status,
            "regulation_refs": refs,
            "verified_count": len(verified),
        })

    return hazards, summaries


# ============== Fusion（异步） ==============
def _build_fusion_prompt(candidate: dict, kb_text: str) -> str:
    path = " > ".join(
        x for x in [candidate.get("category1", ""), candidate.get("category2", ""), candidate.get("category3", "")]
        if x
    )
    return (
        "请对比以下图片：\n"
        "1) 第一张是待检测现场图片（Scene）。\n"
        "2) 后续图片是知识库提供的参考图片（如有）。\n\n"
        f"任务：仅判断三级候选「{candidate.get('issue_id', '')}｜{path}」在现场是否存在。\n"
        "如果不存在，只输出：{\"hazard_count\":0,\"hazards\":[]}。\n"
        "如果存在，只输出严格 JSON，且 hazards 中仅包含这一条隐患：\n"
        "{\n"
        '  "hazard_count": 1,\n'
        '  "hazards": [\n'
        "    {\n"
        f'      "category1": "{candidate.get("category1", "")}",\n'
        f'      "category2": "{candidate.get("category2", "")}",\n'
        f'      "issue_id": "{candidate.get("issue_id", "")}",\n'
        f'      "category3": "{candidate.get("category3", "")}",\n'
        '      "point": "隐患点（简述）",\n'
        '      "description": "隐患说明（结合现场与证据）",\n'
        '      "regulation_refs": [\n'
        '        {"standard": "GB 50016-2014", "clause": "第6.2.9条"}\n'
        "      ],\n"
        '      "violations": []\n'
        "    }\n"
        "  ]\n"
        "}\n\n"
        "【重要硬性约束】\n"
        "1) 不得换成其他 issue_id 或其他三级问题。\n"
        "2) regulation_refs 只填写知识库证据区明确出现的规范名称和条款号。\n"
        "3) 知识库只有条款号时，violations 必须为 []。\n"
        "4) 只有知识库证据区已提供完整原文时，violations 才可填写；content 必须逐字复制原文，"
        "每项格式为 {\"standard\":\"规范名称\",\"clause\":\"第X条\",\"content\":\"原文\"}。\n"
        "5) 只有知识库已提供完整原文时，violations 才可逐字复制原文。\n"
        "6) 只输出 JSON，不要包含任何解释或代码块标记。\n\n"
        "【知识库证据区】\n"
        f"{_clean_kb_text(kb_text)}"
    )


async def _doubao_fusion(
    client: httpx.AsyncClient,
    sem: asyncio.Semaphore,
    scene_url: str,
    kb_text: str,
    kb_image_url: str,
    candidate: dict,
) -> dict:
    image_urls = [scene_url]
    if kb_image_url:
        image_urls.append(kb_image_url)

    payload = _build_responses_payload(
        prompt_text=_build_fusion_prompt(candidate, kb_text),
        image_urls=image_urls,
    )
    headers = {"Content-Type": "application/json", "Authorization": f"Bearer {DOUBAO_API_KEY}"}
    async with sem:
        r = await client.post(DOUBAO_API_URL, headers=headers, json=payload)
    if r.status_code != 200:
        return {"present": False, "hazard": None, "raw": "", "reason": f"http_{r.status_code}"}

    raw = _extract_response_text(r.json())
    try:
        fusion = json.loads(_extract_json(raw))
    except Exception:
        return {"present": False, "hazard": None, "raw": raw, "reason": "parse_error"}

    hazards = (fusion or {}).get("hazards") or []
    if (fusion or {}).get("hazard_count", 0) < 1 or not hazards:
        return {"present": False, "hazard": None, "raw": raw, "reason": "model_no_hazard"}

    hz = hazards[0]
    # 分类字段必须以问题分级表为准，不采用模型自由生成的分类名称。
    for key in [
        "category1_code", "category1", "category2_code", "category2",
        "issue_id", "category3", "regulation_index_raw",
    ]:
        hz[key] = candidate.get(key, "")

    hz["regulation_refs"] = _prepare_regulation_refs(hz, kb_text)
    hz["violations"] = _verified_violations_from_kb(hz, kb_text)
    hz["review_required"] = bool(hz["regulation_refs"])
    return {"present": True, "hazard": hz, "raw": raw}


# 三轮 Fusion 中只要任意一轮判定存在，即保留该隐患。
async def ensemble_fusion_for_type_async(
    scene_url: str,
    kb_text: str,
    kb_image_url: str,
    candidate: dict,
    rounds: int,
    model_sem: asyncio.Semaphore,
    client: httpx.AsyncClient,
):
    tasks = [
        _doubao_fusion(client, model_sem, scene_url, kb_text, kb_image_url, candidate)
        for _ in range(rounds)
    ]
    outputs = await asyncio.gather(*tasks, return_exceptions=True)
    votes_present = 0
    accepted: List[dict] = []

    for i, output in enumerate(outputs, 1):
        if isinstance(output, Exception):
            detail = getattr(output, "detail", None)
            status_code = getattr(output, "status_code", None)
            print(f"[Fusion] {candidate['issue_id']} 第{i}轮异常：{output}" f"type={type(output).__name__}, " f"status_code={status_code}, "
        f"detail={detail}, "
        f"repr={output!r}")
            continue
        if output.get("present") and output.get("hazard"):
            votes_present += 1
            accepted.append(output["hazard"])

    threshold = 1
    decided = votes_present >= threshold
    best_hazard = None
    if decided and accepted:
        best_hazard = max(
            accepted,
            key=lambda h: (
                len(h.get("regulation_refs", [])),
                len(h.get("violations", [])),
                len(h.get("description", "")),
            ),
        )

    return {
        "category1_code": candidate.get("category1_code", ""),
        "category1": candidate.get("category1", ""),
        "category2_code": candidate.get("category2_code", ""),
        "category2": candidate.get("category2", ""),
        "issue_id": candidate.get("issue_id", ""),
        "category3": candidate.get("category3", ""),
        "present_count": votes_present,
        "rounds": rounds,
        "threshold": threshold,
        "present": decided,
        "hazard": best_hazard,
    }


# ============== 主接口（异步并发） ==============
@app.post("/analyze/fire-safety/guess-types")
async def guess_hazard_types_and_kb(image: UploadFile = File(...)):
    """
    简化流程：
    - Stage0：判断是否属于建筑消防巡检场景；
    - StageA：直接从问题分级表输出5个一级/二级/三级候选；
    - KB：按 issue_id 和完整三级路径检索案例；
    - Fusion：每个候选运行3轮，任意1轮判定存在即保留；
    - Clause KB：根据案例库条文号生成 record_id，精确查询规范原文。
    """
    key = upload_image_to_tos(image)
    url = _presigned_get_url(key, expires=600)

    catalog, allowed_text = _get_kb_categories()

    timeout = httpx.Timeout(FS_HTTP_TIMEOUT)
    async with httpx.AsyncClient(timeout=timeout) as client:
        model_sem = asyncio.Semaphore(FS_MODEL_CONCURRENCY)
        kb_sem = asyncio.Semaphore(FS_KB_CONCURRENCY)

        stage0_result = {"relevant": True,"confidence": 1.0,"scene_tags": [],"reason": "disabled",}
        if FS_STAGE0_ENABLED:
            stage0_result = await run_stage0_guard_async(client, model_sem, url)

        debug_log(f"[Stage0] result = {stage0_result}")

        if FS_STAGE0_ENABLED and (not stage0_result.get("relevant", False)
            or float(stage0_result.get("confidence", 0.0)) < FS_STAGE0_THRESHOLD
        ):
            debug_log("[Stage0] 非建筑/消防相关或置信度不足，短路返回空结果。")
            return JSONResponse({
                "image_key": key,
                "image_url": url,
                "stage0": stage0_result,
                "hazard_types": [],
                "kb_results": [],
                "clause_kb_results": [],
                "final_result": {
                    "analysis_status": "stage0_rejected",
                    "catalog_match_status": "not_applicable",
                    "hazard_count": 0,
                    "hazards": [],
                    "unmatched_suggestions": [],
                    "review_required": False,
                    "review_reason": "图片未通过建筑消防巡检场景筛查，未执行后续隐患分析",
                    "analysis_message": "图片未通过建筑消防巡检场景筛查。",
                },
                "type_votes": [],
                "fusion_votes": [],
            })

        selected_types, type_votes, per_round = await ensemble_stageA_async(
            url,allowed_text,catalog,
            rounds=FS_STAGEA_ROUNDS,temps=FS_TEMP_LIST,
            model_sem=model_sem,client=client,
        )

        debug_log("[StageA] 进入后续处理的一级/二级/三级候选：")
        for index, candidate in enumerate(selected_types, 1):
            debug_log(
                f"  {index}. {candidate['category1']} - {candidate['category2']} - "
                f"{candidate['issue_id']} {candidate['category3']}"
            )

        # Stage0已经通过，但StageA三轮均未留下有效白名单候选：
        # 不再直接等同于“无隐患”，而是执行一次开放式模型复核。
        if not selected_types:
            fallback_result = await run_catalog_unmatched_fallback_async(
                client,
                model_sem,
                url,
            )
            final_result = build_catalog_unmatched_final_result(fallback_result)
            return JSONResponse({
                "image_key": key,
                "image_url": url,
                "stage0": stage0_result,
                "hazard_types": [],
                "kb_results": [],
                "clause_kb_results": [],
                "stageA_fallback": fallback_result,
                "final_result": final_result,
                "type_votes": type_votes,
                "fusion_votes": [],
            })

        kb_results = await search_kb_for_types_async(
            selected_types,topk=KB_PER_TYPE_TOPK,kb_sem=kb_sem,client=client,
        )

        debug_log("[StageB] KB 命中摘要：")
        for item in kb_results:
            hit = (
                item["results"][0]["content"][:120].replace("\n", " ") + "..."
                if item.get("results") else "(无结果)"
            )
            debug_log(f"  - {item['issue_id']} {item['category3']} => {hit}")

        fusion_tasks = []
        for candidate in selected_types:
            kb_pack = next(
                (item for item in kb_results if item.get("issue_id") == candidate.get("issue_id")),
                None,
            )
            if not kb_pack or not kb_pack.get("results"):
                print(f"[Fusion] 跳过（无KB证据）：{candidate['issue_id']} {candidate['category3']}")

                async def _noop(candidate=candidate):
                    return {
                        "category1_code": candidate.get("category1_code", ""),
                        "category1": candidate.get("category1", ""),
                        "category2_code": candidate.get("category2_code", ""),
                        "category2": candidate.get("category2", ""),
                        "issue_id": candidate.get("issue_id", ""),
                        "category3": candidate.get("category3", ""),
                        "present_count": 0,
                        "rounds": FS_FUSION_ROUNDS,
                        "threshold": 1,
                        "present": False,
                        "hazard": None,
                    }

                fusion_tasks.append(_noop())
                continue

            kb_text = kb_pack["results"][0].get("content", "") or ""
            kb_image = (kb_pack["results"][0].get("images") or [None])[0]
            fusion_tasks.append(
                ensemble_fusion_for_type_async(
                    url,kb_text,kb_image,candidate,
                    rounds=FS_FUSION_ROUNDS,model_sem=model_sem,client=client,
                )
            )

        fusion_results = await asyncio.gather(*fusion_tasks, return_exceptions=True)

    final_hazards = []
    fusion_votes = []
    for result in fusion_results:
        if isinstance(result, Exception):
            print(f"[Fusion] 任务异常：{result}")
            continue

        fusion_votes.append({
            "category1": result["category1"],
            "category2": result["category2"],
            "issue_id": result["issue_id"],
            "category3": result["category3"],
            "present_count": result["present_count"],
            "rounds": result["rounds"],
            "threshold": result["threshold"],
            "present": result["present"],
        })

        status = "存在(至少1票)" if result["present"] else "不存在(0票)"
        print(
            f"[StageC] {result['issue_id']} {result['category3']} => "
            f"{status} {result['present_count']}/{result['rounds']}"
        )
        if result["present"] and result["hazard"]:
            final_hazards.append(result["hazard"])

    clause_kb_results = []
    if final_hazards:
        async with httpx.AsyncClient(timeout=httpx.Timeout(FS_HTTP_TIMEOUT)) as clause_client:
            clause_sem = asyncio.Semaphore(CLAUSE_KB_CONCURRENCY)
            final_hazards, clause_kb_results = await enrich_hazards_with_clause_kb_async(
                final_hazards,
                clause_sem,
                clause_client,
            )

    review_required = any(
        hazard.get("clause_kb_status") != "kb_exact_match"
        for hazard in final_hazards
    )
    final_result = {
        "analysis_status": "hazard_detected" if final_hazards else "no_hazard",
        "catalog_match_status": "matched" if selected_types else "no_stageA_match",
        "hazard_count": len(final_hazards),
        "hazards": final_hazards,
        "unmatched_suggestions": [],
        "review_required": review_required,
        "review_reason": (
            "部分条文原文未精确匹配，需专业人员复核"
            if review_required
            else ""
        ),
        "analysis_message": (
            "已识别到问题分级表内的疑似隐患。"
            if final_hazards
            else "未识别到最终保留隐患。"
        ),
    }

    return JSONResponse({
        "image_key": key,
        "image_url": url,
        "stage0": stage0_result,
        # 直接输出一级、二级、三级候选。
        "hazard_types": selected_types,
        "kb_results": kb_results,
        "clause_kb_results": clause_kb_results,
        "final_result": final_result,
        "type_votes": type_votes,
        "fusion_votes": fusion_votes,
    })


# （可选）本地直接运行
if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8000")))
