import os
import uuid
import hashlib
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urlparse
from io import BytesIO

import requests
import boto3
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from fastapi import (
    FastAPI,
    File,
    UploadFile,
    HTTPException,
    Depends,
    Form,
    Query,
    Body,
)
from fastapi.responses import JSONResponse, StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from pymongo import MongoClient, DESCENDING, ReturnDocument
from pymongo.errors import DuplicateKeyError, PyMongoError
from dotenv import load_dotenv

from auth import verify_token

# ==== 从 app.py 引用新流水线的配置和函数 ====
import asyncio
import httpx
from app import (
    FS_HTTP_TIMEOUT,
    FS_STAGE0_ENABLED,
    FS_STAGE0_THRESHOLD,
    FS_STAGEA_ROUNDS,
    FS_TEMP_LIST,
    FS_FUSION_ROUNDS,
    FS_MODEL_CONCURRENCY,
    FS_KB_CONCURRENCY,
    KB_PER_TYPE_TOPK,
    _get_kb_categories,
    run_stage0_guard_async,
    ensemble_stageA_async,
    search_kb_for_types_async,
    ensemble_fusion_for_type_async,
    collect_fusion_results_and_enrich_clause_kb_async,
    build_skipped_fusion_result_async,
    build_stage0_rejected_pipeline_payload,
    build_catalog_unmatched_pipeline_payload,
    build_matched_final_result,
    build_completed_pipeline_payload,
    run_catalog_unmatched_fallback_async,
    set_trace_context,
    reset_trace_context,
    trace_log,
    set_clause_cache_collection,
)

# ================== 初始化 ==================
load_dotenv()
VERBOSE_LOG = os.getenv("VERBOSE_LOG", "0") == "1"
FS_BATCH_UPLOAD_ENABLED = os.getenv("FS_BATCH_UPLOAD_ENABLED", "0") == "1"
FS_BATCH_MAX_IMAGES = max(1, int(os.getenv("FS_BATCH_MAX_IMAGES", "20")))
FS_IMAGE_JOB_CONCURRENCY = max(
    1, int(os.getenv("FS_IMAGE_JOB_CONCURRENCY", "3"))
)

analysis_executor = ThreadPoolExecutor(
    max_workers=FS_IMAGE_JOB_CONCURRENCY,
    thread_name_prefix="fire-analysis",
)

def debug_log(*args, **kwargs):
    if VERBOSE_LOG:
        print(*args, **kwargs)

app = FastAPI()


@app.on_event("shutdown")
def shutdown_analysis_executor() -> None:
    analysis_executor.shutdown(wait=False, cancel_futures=False)

# CORS 目前允许所有前端地址访问接口
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# === TOS 配置 ===
TOS_ENDPOINT = "tos-cn-beijing.volces.com"
TOS_BUCKET = "nanjing-fire"
TOS_REGION = "cn-beijing"
TOS_ACCESS_KEY = os.getenv("TOS_ACCESS_KEY")
TOS_SECRET_KEY = os.getenv("TOS_SECRET_KEY")


# === MongoDB 配置 ===
MONGO_URI = os.getenv("MONGODB_URI")
mongo_client = MongoClient(MONGO_URI)
db = mongo_client["fire_safety"]
collection = db["analysis_results"]
clause_cache_collection = db["clause_cache"]
set_clause_cache_collection(clause_cache_collection)

# 对 sha256 建唯一索引，避免并发上传同一图片时重复分析。
# 历史数据没有 sha256 字段，不受该部分索引影响。
try:
    # 兼容曾经试运行过旧版去重代码的环境。索引不存在时忽略即可。
    try:
        collection.drop_index("uniq_image_sha256_cache_version")
    except PyMongoError:
        pass

    collection.create_index(
        [("sha256", 1)],
        unique=True,
        partialFilterExpression={"sha256": {"$type": "string"}},
        name="uniq_sha256",
    )
except PyMongoError as exc:
    print(f"[WARN] 创建SHA256去重索引失败：{exc}")

# === 初始化 TOS 客户端 ===
tos_client = boto3.client(
    "s3",
    endpoint_url=f"https://{TOS_ENDPOINT}",
    aws_access_key_id=TOS_ACCESS_KEY,
    aws_secret_access_key=TOS_SECRET_KEY,
    region_name=TOS_REGION,
)

# ================== 工具函数 ==================
def _tos_fixed_host() -> str:
    return f"{TOS_BUCKET}.{TOS_ENDPOINT}"

# 与 app.py 有一份相似但独立的上传实现；对象前缀是 safe_uploads/，不同于 app.py
def _read_upload_bytes(file: UploadFile) -> bytes:
    """一次性读取上传内容，后续同时用于SHA256计算和TOS上传。"""
    data = file.file.read()
    if not data:
        raise HTTPException(status_code=400, detail="上传图片为空")
    return data


def _calculate_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def upload_image_bytes_to_tos(
    image_bytes: bytes,
    filename: str = "",
    content_type: str = "application/octet-stream",
) -> str:
    """上传已读取的图片字节到TOS，避免计算哈希后文件流位于末尾。"""
    ext = os.path.splitext(filename or "")[-1].lower()
    if ext not in [".png", ".jpg", ".jpeg", ".bmp", ".webp"]:
        ext = ".jpg" if not ext else ext

    key = f"safe_uploads/{uuid.uuid4().hex}{ext}"
    debug_log(f"[DEBUG] 上传 Key: {key}")

    try:
        upload_url = tos_client.generate_presigned_url(
            ClientMethod="put_object",
            Params={
                "Bucket": TOS_BUCKET,
                "Key": key,
                "ContentType": content_type or "application/octet-stream",
            },
            ExpiresIn=600,
        )
        parsed = urlparse(upload_url)
        upload_url = upload_url.replace(parsed.netloc, _tos_fixed_host())

        resp = requests.put(
            upload_url,
            data=image_bytes,
            headers={"Content-Type": content_type or "application/octet-stream"},
            timeout=30,
        )
        if resp.status_code != 200:
            print(f"[ERROR] 上传失败: {resp.status_code} {resp.text}")
            raise HTTPException(status_code=500, detail="TOS 上传失败")

        debug_log(f"[DEBUG] 上传成功: {key}")
        return key

    except HTTPException:
        raise
    except Exception as exc:
        print(f"[ERROR] 上传异常: {exc}")
        raise HTTPException(status_code=500, detail="TOS 上传失败")


def upload_image_to_tos(file: UploadFile) -> str:
    """兼容旧调用；新提交接口使用 upload_image_bytes_to_tos。"""
    data = _read_upload_bytes(file)
    return upload_image_bytes_to_tos(
        data,
        filename=file.filename or "",
        content_type=file.content_type or "application/octet-stream",
    )


def _presigned_get_url(key: str, expires: int = 600) -> str:
    """按需生成短期可访问 URL（查询结果时调用）"""
    try:
        url = tos_client.generate_presigned_url(
            ClientMethod="get_object",
            Params={"Bucket": TOS_BUCKET, "Key": key},
            ExpiresIn=expires,
        )
        parsed = urlparse(url)
        return url.replace(parsed.netloc, _tos_fixed_host())
    except Exception as e:
        print(f"[ERROR] 生成GET签名失败: {e}")
        raise HTTPException(status_code=500, detail="生成签名URL失败")

def _merge_category(code: str, name: str) -> str:
    code, name = str(code or "").strip(), str(name or "").strip()
    if not code or name.startswith(code):
        return name or code
    return f"{code}{name}"


def _get_rectification_advice(issue_id: str) -> str:
    """按issue_id从问题卡库回填整改建议，兼容历史MongoDB结果。"""
    issue_id = str(issue_id or "").strip()
    if not issue_id:
        return ""
    try:
        catalog, _ = _get_kb_categories()
        master = (catalog.get("issues_by_id") or {}).get(issue_id) or {}
        return str(master.get("rectification_advice") or "").strip()
    except Exception as exc:
        debug_log(f"[WARN] 整改建议回填失败：issue_id={issue_id}, error={exc}")
        return ""


def _normalize_multiline_text(value: Any) -> str:
    """
    清理接口展示文本中的多余空行。

    规范知识库原文有时会在标题、编号条款之间插入空白行，
    这里仅将连续空行压缩为普通换行，不改动条文文字内容。
    """
    text = str(value or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # 清除每行行尾空格，同时保留条文本身的分行结构。
    lines = [line.rstrip() for line in text.split("\n")]
    text = "\n".join(lines).strip()

    # “\n 空白行 \n”统一压缩成一个换行，避免前端一条空一行。
    text = re.sub(r"\n[ \t]*\n(?:[ \t]*\n)*", "\n", text)
    return text


def _normalize_hazard_output(hazard: dict) -> dict:
    result = dict(hazard)

    # 旧版已废弃字段不再返回。
    result.pop("verified_regulation_text", None)

    # regulation_refs 是规范条款索引、历史结果兼容和报告导出的必要来源。
    # 即使规范知识库没有精确命中原文，也应保留条款索引供调用方核验。
    regulation_refs = result.get("regulation_refs", []) or []
    result["regulation_refs"] = [
        dict(item)
        for item in regulation_refs
        if isinstance(item, dict)
    ]

    # 规范原文去除多余空行。该处理仅影响接口返回展示，
    # 兼容MongoDB中已经保存的历史结果。
    normalized_violations = []
    for item in result.get("violations", []) or []:
        if not isinstance(item, dict):
            continue
        violation = dict(item)
        # 兼容短期产生过法规相关度字段的历史记录，不再向接口返回。
        violation.pop("relevance_score", None)
        violation["content"] = _normalize_multiline_text(
            violation.get("content", "")
        )
        normalized_violations.append(violation)
    result["violations"] = normalized_violations

    result["category1"] = _merge_category(
        result.pop("category1_code", ""),
        result.get("category1", ""),
    )
    result["category2"] = _merge_category(
        result.pop("category2_code", ""),
        result.get("category2", ""),
    )

    # 新结果中的confidence由app.py根据StageA/Fusion多轮投票计算。
    # 历史结果若没有该字段则返回null，避免伪造置信度。
    raw_confidence = result.get("confidence")
    if raw_confidence in (None, ""):
        result["confidence"] = None
    else:
        try:
            result["confidence"] = round(
                max(0.0, min(float(raw_confidence), 1.0)),
                4,
            )
        except (TypeError, ValueError):
            result["confidence"] = None

    rectification_advice = str(
        result.get("rectification_advice") or ""
    ).strip()
    if not rectification_advice:
        rectification_advice = _get_rectification_advice(
            result.get("issue_id", "")
        )
    result["rectification_advice"] = rectification_advice

    return result


def _normalize_unmatched_suggestion(suggestion: dict) -> dict:
    """统一开放式建议字段，避免前端收到空值或异常类型。"""
    suggestion = dict(suggestion or {})
    evidence = suggestion.get("visual_evidence", []) or []
    if not isinstance(evidence, list):
        evidence = [str(evidence)]

    return {
        "suggested_issue": str(
            suggestion.get("suggested_issue") or "未分类疑似问题"
        ).strip(),
        "point": str(suggestion.get("point") or "").strip(),
        "description": str(suggestion.get("description") or "").strip(),
        "visual_evidence": [
            str(item).strip()
            for item in evidence
            if str(item).strip()
        ][:5],
        "recommended_action": str(
            suggestion.get("recommended_action") or ""
        ).strip(),
        "catalog_match_status": "unmatched",
        "review_required": True,
    }


def _normalize_final_result(result: Optional[dict]) -> Optional[dict]:
    """
    统一MongoDB和查询接口中的结果结构。

    前端不能只根据 hazard_count 判断结果，应优先读取 analysis_status：
    - hazard_detected：已匹配问题分级表并保留正式隐患；
    - catalog_unmatched：疑似存在问题，但是没有在问题卡库中匹配到对应的三级问题；
    - no_hazard：未发现有充分证据支持的隐患；
    - stage0_rejected：图片未通过消防巡检场景筛查，例如风景、自拍、食品等。
    """
    if result is None:
        return None

    normalized = dict(result)
    normalized["hazards"] = [
        _normalize_hazard_output(hazard)
        for hazard in normalized.get("hazards", []) or []
        if isinstance(hazard, dict)
    ]
    normalized["unmatched_suggestions"] = [
        _normalize_unmatched_suggestion(item)
        for item in normalized.get("unmatched_suggestions", []) or []
        if isinstance(item, dict)
    ]

    status = str(normalized.get("analysis_status") or "").strip()
    if not status:
        if normalized["hazards"]:
            status = "hazard_detected"
        elif normalized["unmatched_suggestions"]:
            status = "catalog_unmatched"
        else:
            status = "no_hazard"
    normalized["analysis_status"] = status

    normalized["hazard_count"] = len(normalized["hazards"])

    default_messages = {
        "hazard_detected": "已识别到问题分级表内的疑似隐患。",
        "catalog_unmatched": "未在问题分级表里匹配到对应问题，以下为模型建议项。",
        "no_hazard": "未发现有充分图像证据支持的消防隐患。",
        "stage0_rejected": "图片未通过建筑消防巡检场景筛查。",
    }
    normalized["analysis_message"] = (
        str(normalized.get("analysis_message") or "").strip()
        or default_messages.get(status, "分析已完成。")
    )

    if status == "catalog_unmatched":
        normalized["catalog_match_status"] = "unmatched"
        normalized["review_required"] = True
        normalized["review_reason"] = (
            str(normalized.get("review_reason") or "").strip()
            or "疑似问题未匹配到问题分级表，请人工复核并考虑补充问题分级表。"
        )
    elif status == "stage0_rejected":
        normalized["catalog_match_status"] = "not_applicable"
        normalized["review_required"] = False
        normalized["review_reason"] = (
            str(normalized.get("review_reason") or "").strip()
            or "图片未通过建筑消防巡检场景筛查，未执行后续隐患分析"
        )
    elif status == "hazard_detected":
        normalized["catalog_match_status"] = (
            normalized.get("catalog_match_status") or "matched"
        )
        normalized["review_required"] = bool(
            normalized.get("review_required", False)
        )
        normalized["review_reason"] = str(
            normalized.get("review_reason") or ""
        ).strip()
    else:
        normalized["catalog_match_status"] = (
            normalized.get("catalog_match_status") or "no_stageA_match"
        )
        normalized["review_required"] = bool(
            normalized.get("review_required", False)
        )
        normalized["review_reason"] = str(
            normalized.get("review_reason") or ""
        ).strip()

    return normalized


def _mobile_category2_id(category2_code: Any, category2: Any) -> str:
    """为移动端问题分类生成稳定的展示标识，不写入MongoDB。"""
    code = str(category2_code or "").strip()
    name = str(category2 or "").strip()
    return f"category2:{code or name}"


def _mobile_category2_label(category2_code: Any, category2: Any) -> str:
    code = str(category2_code or "").strip()
    name = str(category2 or "").strip()
    if not code or name.startswith(code):
        return name or code
    return f"{code} {name}"


def _mobile_confidence_percent(confidence: Optional[float]) -> Optional[int]:
    if confidence is None:
        return None
    return int(round(confidence * 100))


def _mobile_regulation_key(standard: Any, clause: Any) -> Tuple[str, str]:
    """仅用于当前三级问题内部的法规去重和原文补充。"""
    standard_text = str(standard or "").strip()
    clause_text = str(clause or "").strip()
    standard_code = re.search(
        r"(?:GB|JGJ|XF|GA|CECS|DB\d*)\s*(?:/\s*T)?\s*\d{2,6}(?:\.\d+)?\s*[-—–－−]\s*\d{4}",
        standard_text,
        flags=re.IGNORECASE,
    )
    standard_key = standard_code.group(0) if standard_code else standard_text
    standard_key = re.sub(r"[\W_]+", "", standard_key).lower()
    clause_key = re.sub(r"[\W_]+", "", clause_text).lower()
    return standard_key, clause_key


def _mobile_regulation_title(standard: Any, clause: Any) -> str:
    standard_text = str(standard or "").strip()
    clause_text = str(clause or "").strip()
    if not clause_text:
        return standard_text
    if clause_text.startswith("第"):
        return f"{standard_text} {clause_text}".strip()
    return f"{standard_text} 第{clause_text}条".strip()


def _mobile_category2_options(catalog: dict) -> List[dict]:
    """从《问题分级》返回完整二级分类值集。"""
    options: List[dict] = []
    for level1 in (catalog.get("hierarchy") or {}).values():
        for level2 in (level1.get("category2") or {}).values():
            category2_code = str(level2.get("category2_code") or "").strip()
            category2 = str(level2.get("category2") or "").strip()
            options.append({
                "category2_id": _mobile_category2_id(category2_code, category2),
                "category2_code": category2_code,
                "label": _mobile_category2_label(category2_code, category2),
            })
    return options


def _mobile_hazard_entries(normalized_result: dict, catalog: dict) -> List[dict]:
    """将标准化三级问题转换为移动端候选项；不按二级分类合并详情。"""
    issues_by_id = catalog.get("issues_by_id") or {}
    entries: List[dict] = []
    for index, hazard in enumerate(normalized_result.get("hazards", []) or []):
        if not isinstance(hazard, dict):
            continue
        issue_id = str(hazard.get("issue_id") or "").strip()
        master = issues_by_id.get(issue_id) or {}
        category2_code = str(master.get("category2_code") or "").strip()
        category2_name = str(master.get("category2") or hazard.get("category2") or "").strip()
        confidence = hazard.get("confidence")
        entries.append({
            "hazard_id": f"h-{index}",
            "source_index": index,
            "raw_hazard": hazard,
            "category2_id": _mobile_category2_id(category2_code, category2_name),
            "category2_code": category2_code,
            "category2": _mobile_category2_label(category2_code, category2_name),
            "issue_id": issue_id,
            "category3": str(hazard.get("category3") or master.get("category3") or "").strip(),
            "confidence": confidence,
        })
    return entries


def _mobile_regulations(hazard: dict) -> List[dict]:
    """只构建当前三级问题的法规列表，绝不读取其他hazard的数据。"""
    verified_by_key: Dict[Tuple[str, str], dict] = {}
    for violation in hazard.get("violations", []) or []:
        if not isinstance(violation, dict):
            continue
        standard = violation.get("standard") or violation.get("standard_title") or ""
        clause = violation.get("clause") or violation.get("clause_no") or ""
        key = _mobile_regulation_key(standard, clause)
        if not all(key):
            continue
        if key not in verified_by_key:
            verified_by_key[key] = violation

    regulations: List[dict] = []
    seen = set()
    for reference in hazard.get("regulation_refs", []) or []:
        if not isinstance(reference, dict):
            continue
        standard = (
            reference.get("standard_title")
            or reference.get("standard")
            or reference.get("standard_code")
            or ""
        )
        clause = (
            reference.get("requested_clause_no")
            or reference.get("clause_no")
            or reference.get("clause")
            or ""
        )
        key = _mobile_regulation_key(standard, clause)
        if not all(key) or key in seen:
            continue
        seen.add(key)
        verified = verified_by_key.get(key) or {}
        regulations.append({
            "standard": str(standard).strip(),
            "clause": str(clause).strip(),
            "title": _mobile_regulation_title(standard, clause),
            "content": _normalize_multiline_text(verified.get("content", "")),
            "verification_status": str(
                verified.get("verification_status")
                or ("已核对原文" if verified.get("content") else "")
                or reference.get("verification_status")
                or "待专业核对"
            ).strip(),
            "source": str(reference.get("source") or "条款索引").strip(),
        })

    # 少数历史记录只有已核验原文；仍然仅追加当前问题自己的条文。
    for key, violation in verified_by_key.items():
        if key in seen:
            continue
        seen.add(key)
        standard = violation.get("standard") or violation.get("standard_title") or ""
        clause = violation.get("clause") or violation.get("clause_no") or ""
        regulations.append({
            "standard": str(standard).strip(),
            "clause": str(clause).strip(),
            "title": _mobile_regulation_title(standard, clause),
            "content": _normalize_multiline_text(violation.get("content", "")),
            "verification_status": str(
                violation.get("verification_status") or "已核对原文"
            ).strip(),
            "source": "规范知识库原文",
        })

    for index, regulation in enumerate(regulations, 1):
        regulation["regulation_id"] = f"reg-{index}"
        regulation["relevance_rank"] = index
        regulation["default_visible"] = index <= 3
    return regulations


def _mobile_hazard_detail(entry: dict) -> dict:
    """返回一个三级问题的独立详情，供移动端切换后刷新下方卡片。"""
    hazard = entry["raw_hazard"]
    regulations = _mobile_regulations(hazard)
    return {
        "hazard_id": entry["hazard_id"],
        "category2_id": entry["category2_id"],
        "category2": entry["category2"],
        "issue_id": entry["issue_id"],
        "category3": entry["category3"],
        "point": str(hazard.get("point") or "").strip(),
        "description": str(hazard.get("description") or "").strip(),
        "confidence": entry["confidence"],
        "confidence_percent": _mobile_confidence_percent(entry["confidence"]),
        "rectification_advice": str(hazard.get("rectification_advice") or "").strip(),
        "regulations": regulations,
        "default_regulation_ids": [
            item["regulation_id"] for item in regulations if item["default_visible"]
        ],
        "remaining_regulation_ids": [
            item["regulation_id"] for item in regulations if not item["default_visible"]
        ],
    }


def _build_mobile_view(normalized_result: dict) -> dict:
    """构建移动端首页展示模型，所有展示状态均为接口临时数据。"""
    catalog, _ = _get_kb_categories()
    entries = _mobile_hazard_entries(normalized_result, catalog)
    groups: Dict[str, List[dict]] = {}
    for entry in entries:
        groups.setdefault(entry["category2_id"], []).append(entry)

    recommended = []
    for category2_id, members in groups.items():
        confidences = [item["confidence"] for item in members if item["confidence"] is not None]
        category_confidence = max(confidences) if confidences else None
        first = members[0]
        default_member = sorted(
            members,
            key=lambda item: (
                item["confidence"] is None,
                -(item["confidence"] or 0.0),
                item["source_index"],
            ),
        )[0]
        recommended.append({
            "category2_id": category2_id,
            "category2_code": first["category2_code"],
            "label": first["category2"],
            "confidence": category_confidence,
            "confidence_percent": _mobile_confidence_percent(category_confidence),
            "issue_count": len(members),
            "default_hazard_id": default_member["hazard_id"],
            "_members": members,
        })

    recommended.sort(
        key=lambda item: (
            item["confidence"] is None,
            -(item["confidence"] or 0.0),
            item["label"],
        )
    )
    default_category = recommended[0] if recommended else None
    default_entry = None
    if default_category:
        default_entry = sorted(
            default_category["_members"],
            key=lambda item: (
                item["confidence"] is None,
                -(item["confidence"] or 0.0),
                item["source_index"],
            ),
        )[0]

    issue_options = []
    for category in recommended:
        for entry in sorted(
            category["_members"],
            key=lambda item: (
                item["confidence"] is None,
                -(item["confidence"] or 0.0),
                item["source_index"],
            ),
        ):
            issue_options.append({
                "hazard_id": entry["hazard_id"],
                "category2_id": entry["category2_id"],
                "category2": entry["category2"],
                "issue_id": entry["issue_id"],
                "category3": entry["category3"],
                "confidence": entry["confidence"],
                "confidence_percent": _mobile_confidence_percent(entry["confidence"]),
            })
        category.pop("_members", None)

    return {
        "field_labels": {"category": "问题分类"},
        "category2_options": _mobile_category2_options(catalog),
        "ai_recommended_categories": recommended,
        "default_category2_id": (
            default_category["category2_id"] if default_category else None
        ),
        "default_hazard_id": default_entry["hazard_id"] if default_entry else None,
        "issue_options": issue_options,
        "selected_hazard": (
            _mobile_hazard_detail(default_entry) if default_entry else None
        ),
    }


# ================== Reused submission response ==================
def _build_reused_submission_payload(doc: dict) -> dict:
    """把已有任务转换为提交接口响应，并重新生成可访问的图片URL。"""
    image_key = doc.get("image_key")
    image_url = _presigned_get_url(image_key, expires=600) if image_key else None
    raw_status = str(doc.get("status") or "processing")
    status = raw_status

    payload = {
        "job_id": doc.get("job_id"),
        "image_key": image_key,
        "image_url": image_url,
        "sha256": doc.get("sha256"),
        "duplicate": True,
        "status": status,
    }
    if status == "done":
        payload["result"] = _normalize_final_result(doc.get("result"))
    elif status == "error":
        payload["error"] = doc.get("error")
    return payload

# ================== 新流水线同步封装 ==================
# 先用本文件的 TOS client 获取现场图 URL，再调用从 app.py 导入的 Stage0、StageA、KB、Fusion；
# 最后用 asyncio.run() 以同步形式返回
def run_new_pipeline_sync(image_key: str, trace_id: str = "") -> dict:
    """
    同步封装：给定 TOS image_key，调用定版 app.py 中的完整流水线：

    Stage0 场景筛查
      → StageA 直接输出一级/二级/三级候选
      → 案例知识库检索
      → Fusion 三状态复核，多数正票直接保留，单票直接视觉证据低置信度召回
      → 规范条文知识库按 record_id 精确补充原文
    """
    trace_tokens = set_trace_context(trace_id or image_key[-16:], image_key)
    trace_log(
        "Pipeline",
        "开始完整消防隐患分析",
        {
            "image_key": image_key,
            "stage0_enabled": bool(FS_STAGE0_ENABLED),
            "stage0_threshold": FS_STAGE0_THRESHOLD,
            "stageA_rounds": FS_STAGEA_ROUNDS,
            "fusion_rounds": FS_FUSION_ROUNDS,
        },
    )
    scene_url = _presigned_get_url(image_key, expires=600)
    catalog, allowed_text = _get_kb_categories()
    timeout = httpx.Timeout(FS_HTTP_TIMEOUT)

    async def _run():
        async with httpx.AsyncClient(timeout=timeout) as client:
            model_sem = asyncio.Semaphore(FS_MODEL_CONCURRENCY)
            kb_sem = asyncio.Semaphore(FS_KB_CONCURRENCY)

            # ---------- Stage0 场景筛查 ----------
            stage0_result = {
                "relevant": True,
                "confidence": 1.0,
                "scene_tags": [],
                "reason": "disabled",
            }
            if FS_STAGE0_ENABLED:
                stage0_result = await run_stage0_guard_async(
                    client,model_sem,scene_url,
                )

            trace_log(
                "Pipeline-Stage0",
                "Stage0最终门控判断",
                {
                    "result": stage0_result,
                    "threshold": FS_STAGE0_THRESHOLD,
                    "will_continue": not (
                        FS_STAGE0_ENABLED
                        and (
                            not stage0_result.get("relevant", False)
                            or float(stage0_result.get("confidence", 0.0)) < FS_STAGE0_THRESHOLD
                        )
                    ),
                },
            )

            # 与定版 app.py 保持一致：
            # relevant=false，或 confidence 低于阈值时，停止后续分析。
            if FS_STAGE0_ENABLED and (
                not stage0_result.get("relevant", False)
                or float(stage0_result.get("confidence", 0.0)) < FS_STAGE0_THRESHOLD
            ):
                trace_log("Pipeline-Stop", "Stage0未通过，停止后续分析并返回stage0_rejected")
                return build_stage0_rejected_pipeline_payload(
                    image_key,
                    scene_url,
                    stage0_result,
                )

            # ---------- StageA：直接输出三级候选 ----------
            selected_types, type_votes, per_round = (
                await ensemble_stageA_async(
                    scene_url,
                    allowed_text,
                    catalog,
                    rounds=FS_STAGEA_ROUNDS,
                    temps=FS_TEMP_LIST,
                    model_sem=model_sem,
                    client=client,
                )
            )

            trace_log(
                "Pipeline-StageA",
                "进入知识库与Fusion的候选",
                [
                    {
                        "rank": index,
                        "issue_id": candidate.get("issue_id", ""),
                        "category1": candidate.get("category1", ""),
                        "category2": candidate.get("category2", ""),
                        "category3": candidate.get("category3", ""),
                        "stage_a_score": candidate.get("stage_a_score"),
                        "stage_a_count": candidate.get("stage_a_count", 0),
                        "stage_a_direct_count": candidate.get("stage_a_direct_count", 0),
                    }
                    for index, candidate in enumerate(selected_types, 1)
                ],
            )

            # Stage0通过但StageA没有任何有效白名单候选时，
            # 调用开放式模型复核，区分“确实无隐患”和“白名单未覆盖”。
            if not selected_types:
                trace_log(
                    "Pipeline-StageA",
                    "StageA没有留下任何白名单候选，进入开放式兜底",
                )
                fallback_result = await run_catalog_unmatched_fallback_async(
                    client,
                    model_sem,
                    scene_url,
                )
                return build_catalog_unmatched_pipeline_payload(
                    image_key,
                    scene_url,
                    stage0_result,
                    fallback_result,
                    type_votes,
                )

            # ---------- 案例知识库检索 ----------
            kb_results = await search_kb_for_types_async(
                selected_types,topk=KB_PER_TYPE_TOPK,kb_sem=kb_sem,client=client,)

            debug_log("[StageB] 案例知识库命中摘要：")
            for item in kb_results:
                if item.get("results"):
                    hit = (
                        item["results"][0]
                        .get("content", "")[:120]
                        .replace("\n", " ")
                        + "..."
                    )
                else:
                    hit = "(无结果)"
                debug_log(
                    f"  - {item.get('issue_id', '')} "
                    f"{item.get('category3', '')} => {hit}"
                )

            # ---------- Fusion：按 issue_id 对齐候选与案例 ----------
            fusion_tasks = []
            for candidate in selected_types:
                issue_id = candidate.get("issue_id", "")
                kb_pack = next(
                    (
                        item
                        for item in kb_results
                        if item.get("issue_id") == issue_id
                    ),
                    None,
                )

                if not kb_pack or not kb_pack.get("results"):
                    trace_log(
                        "Pipeline-Fusion-Skip",
                        f"候选{issue_id}因案例知识库无结果而跳过Fusion",
                        {
                            "category3": candidate.get("category3", ""),
                            "consequence": "该候选present_count=0；若所有候选都被跳过，最终会是no_hazard",
                        },
                    )

                    fusion_tasks.append(
                        build_skipped_fusion_result_async(
                            candidate,
                            category_formatter=_merge_category,
                        )
                    )
                    continue

                kb_item = kb_pack["results"][0]
                kb_text = kb_item.get("content", "") or ""
                kb_img = (kb_item.get("images") or [None])[0]

                try:
                    kb_score = max(
                        0.0,
                        min(float(kb_item.get("score", 0.0) or 0.0), 1.0),
                    )
                except (TypeError, ValueError):
                    kb_score = 0.0

                # 下划线字段仅参与内部置信度与保留判断，不写入MongoDB。
                fusion_candidate = dict(candidate)
                fusion_candidate["_kb_score"] = kb_score

                fusion_tasks.append(
                    ensemble_fusion_for_type_async(
                        scene_url,
                        kb_text,
                        kb_img,
                        fusion_candidate,
                        rounds=FS_FUSION_ROUNDS,
                        model_sem=model_sem,
                        client=client,
                    )
                )

            fusion_results = await asyncio.gather(*fusion_tasks,return_exceptions=True,)

            # ---------- 汇总 Fusion，并在同一客户端内查询规范条文知识库 ----------
            # ---------- 规范条文知识库 ----------
            final_hazards, fusion_votes, clause_kb_results = (
                await collect_fusion_results_and_enrich_clause_kb_async(
                    fusion_results,
                    client,
                    log_stage_c=lambda message: debug_log(f"[StageC] {message}"),
                    log_task_error=lambda message: debug_log(f"[Fusion] {message}"),
                )
            )

        final_result = build_matched_final_result(final_hazards)

        trace_log(
            "Pipeline-Final",
            "完整分析结束",
            {
                "analysis_status": final_result.get("analysis_status"),
                "hazard_count": final_result.get("hazard_count", 0),
                "retained_issue_ids": [
                    hazard.get("issue_id", "")
                    for hazard in final_result.get("hazards", []) or []
                ],
                "fusion_votes": fusion_votes,
                "no_hazard_reason_hint": (
                    "所有候选均无有效present票、被知识库无结果跳过、或Fusion调用/解析无效"
                    if not final_hazards
                    else ""
                ),
            },
        )

        # "stageA_rounds": per_round,
        return build_completed_pipeline_payload(
            image_key,
            scene_url,
            stage0_result,
            selected_types,
            kb_results,
            clause_kb_results,
            final_result,
            type_votes,
            fusion_votes,
        )

    # 全局分析线程池在线程中执行该同步函数，
    # 因此这里使用 asyncio.run() 启动单张图片的完整异步流水线。
    try:
        return asyncio.run(_run())
    except Exception as exc:
        trace_log(
            "Pipeline-ERROR",
            "完整分析发生未捕获异常",
            {"type": type(exc).__name__, "error": repr(exc)},
        )
        raise
    finally:
        reset_trace_context(trace_tokens)


# ------------------ 后台任务 ------------------
def _process_job(job_id: str, image_key: str):
    """
    后台异步分析任务：
    - 使用 app.py 中的 Stage0 + StageA + 案例KB + Fusion + 规范条文KB
    - 正常时将 final_result 写入 result 字段
    - 任意异常（超时 / API 调用错误 / 其它异常）：status=error，result=null
    """
    print(f"[JOB] 开始处理 job_id={job_id} image_key={image_key}", flush=True)
    try:
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        # 初始化 / 标记 processing（幂等），顺便清空上一次的 result / error
        collection.update_one(
            {"job_id": job_id},
            {
                "$set": {
                    "status": "processing",
                    "image_key": image_key,
                    "result": None,  # 提交处理时先置空
                    "error": None,
                },
                "$setOnInsert": {"create_time": now_str},
            },
            upsert=True,
        )

        # 调用新流水线
        pipeline_res = run_new_pipeline_sync(image_key, trace_id=job_id)
        final_result = _normalize_final_result(
            pipeline_res.get(
                "final_result",
                {
                    "analysis_status": "no_hazard",
                    "hazard_count": 0,
                    "hazards": [],
                    "unmatched_suggestions": [],
                },
            )
        )

        update_doc = {
            "status": "done",
            "result": final_result,  # 成功时写入正常结果
            "finish_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "error": None,
            # 同时把流水线的中间信息也存起来（可选）
            # "stage0": pipeline_res.get("stage0"),
            # "hazard_types": pipeline_res.get("hazard_types"),
            # "kb_results": pipeline_res.get("kb_results"),
            # "clause_kb_results": pipeline_res.get("clause_kb_results", []),
            # "type_votes": pipeline_res.get("type_votes"),
            # "fusion_votes": pipeline_res.get("fusion_votes"),
        }

        collection.update_one(
            {"job_id": job_id},
            {"$set": update_doc},
            upsert=True,
        )

        print(
            f"[JOB] 完成 job_id={job_id} "
            f"(hazard_count={final_result.get('hazard_count', 0)})"
        )

    except Exception as e:
        # 任意异常（包括超时、API 调用错误等）都走这里
        print(f"[JOB ERROR] job_id={job_id}: {e}")
        collection.update_one(
            {"job_id": job_id},
            {
                "$set": {
                    "status": "error",
                    "result": None,  # 关键：异常时 result 置为 null
                    "error": str(e),
                    "finish_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
            },
            upsert=True,
        )


# ================== 提交与查询接口 ==================
def _enqueue_analysis_job(job_id: str, image_key: str) -> None:
    print(f"[QUEUE] 入队 job_id={job_id}", flush=True)
    try:
        analysis_executor.submit(_process_job, job_id, image_key)
    except Exception as exc:
        collection.update_one(
            {"job_id": job_id},
            {
                "$set": {
                    "status": "error",
                    "result": None,
                    "error": str(exc),
                    "finish_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
            },
        )
        raise


async def _prepare_single_submission(
    image: UploadFile,
    requested_job_id: Optional[str] = None,
) -> dict:
    """
    提交图片：
    - 读取图片字节并计算 SHA256；
    - MongoDB 顶层只新增 sha256 和 duplicate 两个字段；
    - 相同 SHA256 已完成时直接返回原结果；
    - 相同 SHA256 正在处理中时复用原 job_id；
    - 首次图片才上传 TOS 并启动完整分析。
    """
    try:
        image_bytes = _read_upload_bytes(image)
        sha256 = _calculate_sha256(image_bytes)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        existing = collection.find_one({"sha256": sha256})
        if existing:
            existing_status = str(existing.get("status") or "processing")

            # 只更新 duplicate，不新增其他去重相关字段。
            collection.update_one(
                {"_id": existing["_id"]},
                {"$set": {"duplicate": True}},
            )
            existing["duplicate"] = True

            # 已完成：不上传、不调用模型，直接返回已有结果。
            if existing_status == "done" and existing.get("result") is not None:
                print(
                    f"[DEDUP] 命中已完成旧结果，不会重新运行流水线："
                    f"sha256={sha256} job_id={existing.get('job_id')} status=done",
                    flush=True,
                )
                return _build_reused_submission_payload(existing)

            # 正在上传或分析：复用原任务，前端继续轮询原 job_id。
            if existing_status in {"uploading", "queued", "processing"}:
                print(
                    f"[DEDUP] 命中正在处理的旧任务："
                    f"sha256={sha256} job_id={existing.get('job_id')} status={existing_status}",
                    flush=True,
                )
                return _build_reused_submission_payload(existing)

            # 上次失败：通过 status 原子抢占重试权，不增加重试次数等字段。
            claimed = collection.find_one_and_update(
                {"_id": existing["_id"], "status": "error"},
                {
                    "$set": {
                        "status": "uploading",
                        "result": None,
                        "error": None,
                        "duplicate": True,
                    }
                },
                return_document=ReturnDocument.AFTER,
            )
            if claimed:
                image_key = claimed.get("image_key")
                if not image_key:
                    try:
                        image_key = upload_image_bytes_to_tos(
                            image_bytes,
                            filename=image.filename or "",
                            content_type=image.content_type or "application/octet-stream",
                        )
                    except Exception as exc:
                        collection.update_one(
                            {"_id": claimed["_id"]},
                            {
                                "$set": {
                                    "status": "error",
                                    "result": None,
                                    "error": str(exc),
                                    "finish_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                }
                            },
                        )
                        raise

                    collection.update_one(
                        {"_id": claimed["_id"]},
                        {"$set": {"image_key": image_key}},
                    )
                    claimed["image_key"] = image_key

                collection.update_one(
                    {"_id": claimed["_id"]},
                    {"$set": {"status": "queued", "error": None}},
                )
                claimed["status"] = "queued"
                _enqueue_analysis_job(claimed.get("job_id"), image_key)
                return _build_reused_submission_payload(claimed)

            latest = collection.find_one({"sha256": sha256})
            if latest:
                return _build_reused_submission_payload(latest)

        # 首次图片：插入占位记录。sha256 唯一索引防止并发重复分析。
        job_id = requested_job_id or uuid.uuid4().hex
        if collection.find_one({"job_id": job_id}):
            raise HTTPException(status_code=409, detail="job_id 已存在，请更换 job_id")

        placeholder = {
            "job_id": job_id,
            "status": "uploading",
            "image_key": None,
            "sha256": sha256,
            "duplicate": False,
            "result": None,
            "error": None,
            "create_time": now,
        }

        try:
            insert_result = collection.insert_one(placeholder)
        except DuplicateKeyError:
            # 几乎同时上传同一张图时，后到请求复用先到任务。
            existing = collection.find_one({"sha256": sha256})
            if existing:
                collection.update_one(
                    {"_id": existing["_id"]},
                    {"$set": {"duplicate": True}},
                )
                existing["duplicate"] = True
                return _build_reused_submission_payload(existing)
            raise

        try:
            image_key = upload_image_bytes_to_tos(
                image_bytes,
                filename=image.filename or "",
                content_type=image.content_type or "application/octet-stream",
            )
        except Exception as exc:
            collection.update_one(
                {"_id": insert_result.inserted_id},
                {
                    "$set": {
                        "status": "error",
                        "result": None,
                        "error": str(exc),
                        "finish_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    }
                },
            )
            raise

        collection.update_one(
            {"_id": insert_result.inserted_id},
            {
                "$set": {
                    "status": "queued",
                    "image_key": image_key,
                    "error": None,
                }
            },
        )

        _enqueue_analysis_job(job_id, image_key)

        preview_url = _presigned_get_url(image_key, expires=600)
        return {
            "job_id": job_id,
            "image_key": image_key,
            "image_url": preview_url,
            "sha256": sha256,
            "duplicate": False,
            "status": "queued",
        }

    except HTTPException:
        raise
    except Exception as exc:
        print(f"[ERROR] 提交失败: {exc}")
        raise HTTPException(status_code=500, detail="提交失败")



@app.post("/analyze/fire-safety", dependencies=[Depends(verify_token)])
async def submit_analysis(
    image: UploadFile = File(...),
    job_id: Optional[str] = Form(None),
):
    payload = await _prepare_single_submission(image, requested_job_id=job_id)
    return JSONResponse(payload)


_BATCH_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}


@app.post("/analyze/fire-safety/batch", dependencies=[Depends(verify_token)])
async def submit_analysis_batch(
    images: Optional[List[UploadFile]] = File(None),
):
    """批量提交图片；每张图片复用单图接口的既有处理、去重和任务创建逻辑。"""
    if not FS_BATCH_UPLOAD_ENABLED:
        raise HTTPException(
            status_code=503,
            detail="批量上传接口当前未启用，请使用单图接口 /analyze/fire-safety。",
        )

    if not images:
        raise HTTPException(status_code=400, detail="请至少上传1张图片")
    if len(images) > FS_BATCH_MAX_IMAGES:
        raise HTTPException(
            status_code=400,
            detail=f"一次最多上传{FS_BATCH_MAX_IMAGES}张图片",
        )

    items = []
    accepted_count = 0
    failed_count = 0
    for index, image in enumerate(images):
        filename = image.filename or ""
        extension = os.path.splitext(filename)[1].lower()
        if extension not in _BATCH_IMAGE_EXTENSIONS:
            failed_count += 1
            items.append(
                {
                    "index": index,
                    "filename": filename,
                    "status": "rejected",
                    "error": "不支持的图片格式",
                }
            )
            continue

        try:
            payload = await _prepare_single_submission(image)
            accepted_count += 1
            items.append({"index": index, "filename": filename, **payload})
        except HTTPException as exc:
            failed_count += 1
            items.append(
                {
                    "index": index,
                    "filename": filename,
                    "status": "rejected",
                    "error": str(exc.detail),
                }
            )
        except Exception as exc:
            failed_count += 1
            items.append(
                {
                    "index": index,
                    "filename": filename,
                    "status": "rejected",
                    "error": str(exc),
                }
            )

    return JSONResponse(
        {
            "batch_count": len(images),
            "accepted_count": accepted_count,
            "failed_count": failed_count,
            "items": items,
        }
    )


@app.post(
    "/analyze/fire-safety/batch/status",
    dependencies=[Depends(verify_token)],
)
async def get_batch_analysis_status(
    job_ids: List[str] = Body(..., embed=True),
):
    if not job_ids:
        raise HTTPException(status_code=400, detail="job_ids不能为空")

    unique_job_ids = list(dict.fromkeys(str(job_id) for job_id in job_ids))
    docs = collection.find({"job_id": {"$in": unique_job_ids}})
    docs_by_job_id = {str(doc.get("job_id")): doc for doc in docs}

    items = []
    for job_id in job_ids:
        normalized_job_id = str(job_id)
        doc = docs_by_job_id.get(normalized_job_id)
        if not doc:
            items.append(
                {
                    "job_id": normalized_job_id,
                    "status": "not_found",
                    "result": None,
                }
            )
            continue

        item = {
            "job_id": normalized_job_id,
            "status": doc.get("status", "queued"),
            "result": _normalize_final_result(doc.get("result")),
        }
        if doc.get("status") == "error":
            item["error"] = doc.get("error")
        items.append(item)

    return JSONResponse({"items": items})


@app.get("/analyze/fire-safety/{job_id}", dependencies=[Depends(verify_token)])
async def get_result(job_id: str):
    """
    查询结果（用于前端轮询）
    - 每次查询自动为 image_key 生成一个新鲜的签名URL（避免过期）
    返回:
      - status: queued | processing | done | error | not_found
      - image_key: 仅存 key
      - image_url: 动态签名的短效URL（预览用）
      - result: done时包含标准化后的原始分析结果；处理中为null
      - (error) error: 错误信息
    """
    url_expires = 600
    doc = collection.find_one({"job_id": job_id})
    if not doc:
        return {"status": "not_found"}

    image_key = doc.get("image_key")
    image_url = _presigned_get_url(image_key, expires=url_expires) if image_key else None

    payload = {
        "job_id": doc.get("job_id", job_id),
        "status": doc.get("status", "queued"),
        "image_key": image_key,
        "image_url": image_url,  # 每次查询都动态生成，保证不过期
        "sha256": doc.get("sha256"),
        "duplicate": bool(doc.get("duplicate", False)),
        "create_time": doc.get("create_time"),
        "finish_time": doc.get("finish_time"),
        "result": _normalize_final_result(
            doc.get("result", None)
        ),  # done时返回标准化原始结果；处理中仍为null
        # "stage0": doc.get("stage0"),
        # "hazard_types": doc.get("hazard_types", []) or [],
        # "clause_kb_results": doc.get(
        #     "clause_kb_results", []
        # ) or [],
        # "type_votes": doc.get("type_votes", []) or [],
        # "fusion_votes": doc.get("fusion_votes", []) or [],
    }
    if doc.get("status") == "error":
        payload["error"] = doc.get("error")

    return JSONResponse(payload)


@app.get(
    "/analyze/fire-safety/{job_id}/mobile-view",
    dependencies=[Depends(verify_token)],
)
async def get_mobile_view(job_id: str):
    """返回移动端“问题分类”页面的初始展示模型，不修改任务结果。"""
    doc = collection.find_one({"job_id": job_id})
    if not doc:
        return JSONResponse({"job_id": job_id, "status": "not_found", "mobile_view": None})

    raw_status = str(doc.get("status") or "processing")
    status = "processing" if raw_status == "uploading" else raw_status
    image_key = doc.get("image_key")
    payload = {
        "job_id": doc.get("job_id", job_id),
        "status": status,
        "image_key": image_key,
        "image_url": (
            _presigned_get_url(image_key, expires=600)
            if image_key
            else None
        ),
        "mobile_view": None,
    }
    if status == "done":
        normalized_result = _normalize_final_result(doc.get("result") or {}) or {}
        payload["mobile_view"] = _build_mobile_view(normalized_result)
    elif status == "error":
        payload["error"] = doc.get("error")
    return JSONResponse(payload)


@app.get(
    "/analyze/fire-safety/{job_id}/mobile-view/hazards/{hazard_id}",
    dependencies=[Depends(verify_token)],
)
async def get_mobile_hazard_detail(job_id: str, hazard_id: str):
    """切换三级问题时返回该问题独立的描述、法规和整改建议。"""
    doc = collection.find_one({"job_id": job_id})
    if not doc:
        return JSONResponse({"job_id": job_id, "status": "not_found", "selected_hazard": None})

    raw_status = str(doc.get("status") or "processing")
    status = "processing" if raw_status == "uploading" else raw_status
    if status != "done":
        payload = {
            "job_id": doc.get("job_id", job_id),
            "status": status,
            "selected_hazard": None,
        }
        if status == "error":
            payload["error"] = doc.get("error")
        return JSONResponse(payload)

    normalized_result = _normalize_final_result(doc.get("result") or {}) or {}
    catalog, _ = _get_kb_categories()
    entries = _mobile_hazard_entries(normalized_result, catalog)
    selected = next(
        (entry for entry in entries if entry["hazard_id"] == hazard_id),
        None,
    )
    if selected is None:
        raise HTTPException(status_code=404, detail="未找到指定的三级问题")

    return JSONResponse({
        "job_id": doc.get("job_id", job_id),
        "status": "done",
        "selected_hazard": _mobile_hazard_detail(selected),
    })


# ================== 查询历史记录 ==================
@app.get("/analyze/history/query", dependencies=[Depends(verify_token)])
async def query_analysis_history(
    category1: Optional[List[str]] = Query(
        None, description="一级分类，如：1.1建筑专业",
    ),
    category2: Optional[List[str]] = Query(
        None, description="二级分类，如：1.1.2防火分区和防火构造",
    ),
    issue_id: Optional[List[str]] = Query(
        None, description="三级问题编号，如：1.1.2-13",
    ),
    start_date: Optional[str] = Query(
        None,
        description="起始日期，YYYY-MM-DD",
    ),
    end_date: Optional[str] = Query(
        None, description="结束日期，YYYY-MM-DD",
    ),
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
):
    """
    查询已完成的历史分析记录。

    支持：
    - 一级分类筛选；
    - 二级分类筛选；
    - 三级问题编号筛选；
    - 时间段筛选；
    - 分页。
    """
    try:
        query = {"status": "done"}

        # 时间筛选
        if start_date or end_date:
            query["create_time"] = {}
            if start_date:
                query["create_time"]["$gte"] = f"{start_date} 00:00:00"
            if end_date:
                query["create_time"]["$lte"] =  f"{end_date} 23:59:59"

        # 分类筛选
        and_filters = []
        if category1:
            and_filters.append({
                "result.hazards.category1": {
                    "$in": category1
                }
            })

        if category2:
            and_filters.append({
                "result.hazards.category2": {
                    "$in": category2
                }
            })

        if issue_id:
            and_filters.append({
                "result.hazards.issue_id": {
                    "$in": issue_id
                }
            })

        if and_filters:
            query["$and"] = and_filters

        total = collection.count_documents(query)
        skip = (page - 1) * limit

        cursor = (
            collection.find(query)
            .sort("create_time", -1)
            .skip(skip)
            .limit(limit)
        )

        results = []

        for doc in cursor:
            image_key = doc.get("image_key")
            image_url = ( _presigned_get_url(image_key, expires=900, )
                if image_key
                else None
            )

            result_doc = _normalize_final_result(
                doc.get("result", {}) or {}
            ) or {}
            hazards = result_doc.get("hazards", []) or []

            category_set = set()
            issue_set = set()

            for hazard in hazards:
                category2_value = hazard.get( "category2", "", )
                if category2_value:
                    category_set.add(category2_value)

                hazard_issue_id = hazard.get("issue_id", "", )
                if hazard_issue_id:
                    issue_set.add(hazard_issue_id)

            unmatched_suggestions = (
                result_doc.get("unmatched_suggestions", []) or []
            )
            results.append({
                "id": str(doc.get("_id")),
                "job_id": doc.get("job_id", ""),
                "image_url": image_url,
                "analysis_status": result_doc.get(
                    "analysis_status", "no_hazard"
                ),
                "analysis_message": result_doc.get(
                    "analysis_message", ""
                ),
                "hazard_count": result_doc.get("hazard_count", 0),
                "unmatched_suggestion_count": len(
                    unmatched_suggestions
                ),
                "unmatched_suggestions": unmatched_suggestions,
                "review_required": bool(
                    result_doc.get("review_required", False)
                ),
                "review_reason": result_doc.get(
                    "review_reason", ""
                ),
                "categories": sorted(category_set),
                "issue_ids": sorted(issue_set),
                "create_time": (
                    doc.get("finish_time") or doc.get("create_time")
                ),
            })

        return {
            "total": total,
            "page": page,
            "limit": limit,
            "data": results,
        }

    except Exception as e:
        print(f"[ERROR] 查询历史记录失败: {e}")
        raise HTTPException( status_code=500, detail="查询失败",)


# ================== 批量删除接口 ==================
@app.post("/analyze/fire-safety/delete", dependencies=[Depends(verify_token)])
async def delete_analysis_records(payload: dict = Body(...)):
    """
    批量删除分析记录
    - 接收 JSON: { "job_ids": ["id1", "id2", ...] }
    - 返回: { "deleted_count": n }
    """
    job_ids = payload.get("job_ids", [])
    if not job_ids or not isinstance(job_ids, list):
        raise HTTPException(status_code=400, detail="缺少或格式错误的 job_ids 参数")

    result = collection.delete_many({"job_id": {"$in": job_ids}})
    return {"deleted_count": result.deleted_count}


EXPORT_COLUMNS = [
    "序号", "验收图片", "一级分类", "二级分类", "三级问题",
    "隐患说明", "整改建议", "规范条款索引", "规范原文",
    "规范复核提示", "分析时间",
]

# ================== 批量下载接口 ==================
@app.post("/analyze/fire-safety/download", dependencies=[Depends(verify_token)])
async def download_reports(payload: Optional[dict] = Body(default=None)):
    """导出指定 job_ids；请求体为空或 job_ids 为空时导出全部已完成记录。"""
    try:
        payload = payload or {}
        job_ids = payload.get("job_ids", [])
        if not isinstance(job_ids, list):
            raise HTTPException(status_code=400, detail="job_ids必须是数组")

        query = {"status": "done"}
        if job_ids:
            query["job_id"] = {"$in": job_ids}
        cursor = collection.find(query).sort("create_time", DESCENDING)

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "FireSafetyReports"
        worksheet.append(EXPORT_COLUMNS)

        widths = [6, 22, 18, 24, 42, 55, 55, 45, 70, 45, 20]
        if len(widths) != len(EXPORT_COLUMNS):
            raise ValueError("导出表头数量与列宽数量不一致")
        for index, width in enumerate(widths, 1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width

        worksheet.row_dimensions[1].height = 28
        for column in range(1, len(EXPORT_COLUMNS) + 1):
            cell = worksheet.cell(1, column)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(bold=True)

        row, sequence = 2, 1
        for document in cursor:
            result_doc = _normalize_final_result(
                document.get("result", {}) or {}
            ) or {}
            hazards = list(result_doc.get("hazards", []) or [])

            image_data = None
            image_key = document.get("image_key")
            if image_key:
                try:
                    response = requests.get(_presigned_get_url(image_key, expires=600), timeout=5)
                    response.raise_for_status()
                    content_type = response.headers.get("Content-Type", "").lower()
                    if not content_type.startswith("image/"):
                        raise ValueError(f"TOS返回内容不是图片：{content_type}")
                    image_data = response.content
                except Exception as exc:
                    print(f"[ERROR] 下载图片失败: {exc}")

            if not hazards:
                worksheet.cell(row, 1, sequence)
                if image_data:
                    try:
                        excel_image = XLImage(BytesIO(image_data))
                        excel_image.width = excel_image.height = 80
                        worksheet.add_image(excel_image, f"B{row}")
                    except Exception as exc:
                        print(f"[ERROR] 插入图片失败: {exc}")
                        worksheet.cell(row, 2, "图片加载失败")
                else:
                    worksheet.cell(row, 2, "图片加载失败")

                status = result_doc.get("analysis_status", "no_hazard")
                if status == "stage0_rejected":
                    description = "图片未通过消防巡检场景筛查，未执行后续隐患分析"
                elif status == "catalog_unmatched":
                    suggestion_lines = []
                    for index, suggestion in enumerate(
                        result_doc.get("unmatched_suggestions", []) or [],
                        1,
                    ):
                        evidence = "、".join(
                            suggestion.get("visual_evidence", []) or []
                        )
                        parts = [
                            f"建议项{index}：{suggestion.get('suggested_issue', '')}",
                            f"隐患点：{suggestion.get('point', '')}",
                            f"隐患说明：{suggestion.get('description', '')}",
                        ]
                        if evidence:
                            parts.append(f"现场依据：{evidence}")
                        action = suggestion.get("recommended_action", "")
                        if action:
                            parts.append(f"建议复核：{action}")
                        suggestion_lines.append("\n".join(parts))
                    suggestion_text = "\n\n".join(suggestion_lines)
                    description = "未在问题分级表里匹配到对应问题。"
                    if suggestion_text:
                        description += f"\n\n模型开放式建议：\n{suggestion_text}"
                else:
                    description = (
                        result_doc.get("analysis_message")
                        or "未识别到最终保留隐患"
                    )
                worksheet.cell(row, 6, description)
                worksheet.cell(row, 10, result_doc.get("review_reason", ""))
                worksheet.cell(row, 11, document.get("finish_time", ""))
                worksheet.row_dimensions[row].height = 65
                for column in range(3, len(EXPORT_COLUMNS) + 1):
                    worksheet.cell(row, column).alignment = Alignment(wrap_text=True, vertical="center")
                row += 1
                sequence += 1
                continue

            start_row, end_row = row, row + len(hazards) - 1
            if len(hazards) > 1:
                worksheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
                worksheet.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
            worksheet.cell(start_row, 1, sequence).alignment = Alignment(vertical="center")

            if image_data:
                try:
                    excel_image = XLImage(BytesIO(image_data))
                    excel_image.width = excel_image.height = 80
                    worksheet.add_image(excel_image, f"B{start_row}")
                except Exception as exc:
                    print(f"[ERROR] 插入图片失败: {exc}")
                    worksheet.cell(start_row, 2, "图片加载失败")
            else:
                worksheet.cell(start_row, 2, "图片加载失败")

            for index, hazard in enumerate(hazards):
                current_row = start_row + index
                worksheet.row_dimensions[current_row].height = 80
                worksheet.cell(current_row, 3, _merge_category(hazard.get("category1_code", ""), hazard.get("category1", "")))
                worksheet.cell(current_row, 4, _merge_category(hazard.get("category2_code", ""), hazard.get("category2", "")))

                issue_id = str(hazard.get("issue_id") or "").strip()
                category3 = str(hazard.get("category3") or "").strip()
                worksheet.cell(current_row, 5, f"{issue_id} {category3}".strip())

                point = str(hazard.get("point") or "").strip()
                description = str(hazard.get("description") or "").strip()
                hazard_text = (
                    f"隐患点：{point}\n隐患说明：{description}"
                    if point and description and point != description
                    else description or point
                )
                worksheet.cell(current_row, 6, hazard_text)
                worksheet.cell(
                    current_row,
                    7,
                    str(hazard.get("rectification_advice") or "").strip(),
                )

                reference_texts = []
                for reference in hazard.get("regulation_refs", []) or []:
                    standard = reference.get("standard_code") or reference.get("standard") or ""
                    clause = (
                        reference.get("requested_clause_no")
                        or reference.get("clause_no")
                        or reference.get("clause")
                        or ""
                    )
                    text = f"{standard} {clause}".strip()
                    if text and text not in reference_texts:
                        reference_texts.append(text)
                worksheet.cell(current_row, 8, "\n\n".join(reference_texts))

                clause_texts = [
                    str(item.get("content") or "").strip()
                    for item in hazard.get("violations", []) or []
                    if isinstance(item, dict)
                    and str(item.get("content") or "").strip()
                ]
                worksheet.cell(current_row, 9, "\n\n".join(clause_texts))
                worksheet.cell(current_row, 10, hazard.get("regulation_review_status", ""))
                worksheet.cell(current_row, 11, document.get("finish_time", ""))

            for column in range(3, len(EXPORT_COLUMNS) + 1):
                for current_row in range(start_row, end_row + 1):
                    worksheet.cell(current_row, column).alignment = Alignment(wrap_text=True, vertical="center")

            row = end_row + 1
            sequence += 1

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)

        filename = (
            "fire_safety_reports_"
            f"{datetime.now().strftime('%Y%m%d%H%M%S')}"
            ".xlsx"
        )
        headers = {
            "Content-Disposition": ( f'attachment; filename="{filename}"')
        }
        return StreamingResponse(
            stream,
            media_type=( "application/vnd.openxmlformats-officedocument." "spreadsheetml.sheet" ),
            headers=headers,
        )

    except HTTPException:
        raise
    except Exception as exc:
        print(f"[ERROR] 导出失败: {exc}")
        raise HTTPException( status_code=500,detail="导出失败",)
