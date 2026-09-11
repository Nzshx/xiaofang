import os
import re
import uuid
import json
from datetime import datetime
from typing import Optional, List
from urllib.parse import urlparse
from io import BytesIO

import requests
import boto3
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from fastapi import (
    FastAPI,
    File,
    UploadFile,
    HTTPException,
    Depends,
    Form,
    Query,
    BackgroundTasks,
    Body,
)
from fastapi.responses import JSONResponse, StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from pymongo import MongoClient, DESCENDING
from dotenv import load_dotenv

from auth import verify_token

# ==== 从 app.py 引用新流水线的配置和函数 ====
import asyncio
import httpx
from app import (
    FS_HTTP_TIMEOUT,
    FS_STAGE0_ENABLED,
    FS_STAGE0_THRESHOLD,
    FS_STAGEA_ROUNDS,
    FS_TEMP_LIST,
    FS_FUSION_ROUNDS,
    FS_MODEL_CONCURRENCY,
    FS_KB_CONCURRENCY,
    _get_kb_categories,
    run_stage0_guard_async,
    ensemble_stageA_async,
    search_kb_for_types_async,
    ensemble_fusion_for_type_async,
)

# ================== 初始化 ==================
load_dotenv()
app = FastAPI()

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# === TOS 配置 ===
TOS_ENDPOINT = "tos-cn-beijing.volces.com"
TOS_BUCKET = "nanjing-fire"
TOS_REGION = "cn-beijing"
TOS_ACCESS_KEY = os.getenv("TOS_ACCESS_KEY")
TOS_SECRET_KEY = os.getenv("TOS_SECRET_KEY")

# === 豆包模型配置（这里仅保留，实际分析逻辑用 app.py 的新流水线） ===
DOUBAO_API_URL = "https://ark.cn-beijing.volces.com/api/v3/chat/completions"
DOUBAO_API_KEY = os.getenv("DOUBAO_API_KEY")
DOUBAO_MODEL = "doubao-1-5-thinking-vision-pro-250428"

# === MongoDB 配置 ===
MONGO_URI = os.getenv("MONGODB_URI")
mongo_client = MongoClient(MONGO_URI)
db = mongo_client["fire_safety"]
collection = db["analysis_results"]

# === 初始化 TOS 客户端 ===
tos_client = boto3.client(
    "s3",
    endpoint_url=f"https://{TOS_ENDPOINT}",
    aws_access_key_id=TOS_ACCESS_KEY,
    aws_secret_access_key=TOS_SECRET_KEY,
    region_name=TOS_REGION,
)

# ================== 工具函数 ==================
def _tos_fixed_host() -> str:
    return f"{TOS_BUCKET}.{TOS_ENDPOINT}"


def upload_image_to_tos(file: UploadFile) -> str:
    """
    上传图片到 TOS，返回对象 key（保持原来的 safe_uploads 前缀，兼容历史）
    """
    ext = os.path.splitext(file.filename or "")[-1]
    if not ext:
        ext = ".jpg"
    key = f"safe_uploads/{uuid.uuid4().hex}{ext}"
    print(f"[DEBUG] 上传 Key: {key}")

    try:
        content_type = file.content_type or "application/octet-stream"
        # 生成上传 URL（PUT）
        upload_url = tos_client.generate_presigned_url(
            ClientMethod="put_object",
            Params={"Bucket": TOS_BUCKET, "Key": key, "ContentType": content_type},
            ExpiresIn=600,
        )
        # 修正 host
        parsed = urlparse(upload_url)
        upload_url = upload_url.replace(parsed.netloc, _tos_fixed_host())

        # 执行上传
        resp = requests.put(
            upload_url, data=file.file.read(), headers={"Content-Type": content_type}
        )
        if resp.status_code != 200:
            print(f"[ERROR] 上传失败: {resp.status_code} {resp.text}")
            raise HTTPException(status_code=500, detail="TOS 上传失败")

        print(f"[DEBUG] 上传成功: {key}")
        return key

    except Exception as e:
        print(f"[ERROR] 上传异常: {e}")
        raise HTTPException(status_code=500, detail="TOS 上传失败")


def _presigned_get_url(key: str, expires: int = 600) -> str:
    """按需生成短期可访问 URL（查询结果时调用）"""
    try:
        url = tos_client.generate_presigned_url(
            ClientMethod="get_object",
            Params={"Bucket": TOS_BUCKET, "Key": key},
            ExpiresIn=expires,
        )
        parsed = urlparse(url)
        return url.replace(parsed.netloc, _tos_fixed_host())
    except Exception as e:
        print(f"[ERROR] 生成GET签名失败: {e}")
        raise HTTPException(status_code=500, detail="生成签名URL失败")


def _extract_json(text: str) -> str:
    """兼容 code block 包裹；优先取 ```json ... ``` 内部"""
    m = re.search(r"```json\s*(\{.*?\})\s*```", text, re.DOTALL)
    return m.group(1) if m else text.strip()


# ================== 新流水线同步封装 ==================
def run_new_pipeline_sync(image_key: str) -> dict:
    """
    同步封装：给定 TOS image_key，调用 app.py 中的
    Stage0 + StageA + KB + Fusion 流水线，返回结构为：

    {
        "image_key": ...,
        "image_url": ...,
        "stage0": {...},
        "hazard_types": [...],
        "kb_results": [...],
        "final_result": {"hazard_count": int, "hazards": [...]},
        "type_votes": [...],
        "fusion_votes": [...]
    }
    """
    scene_url = _presigned_get_url(image_key, expires=600)
    mapping, allowed_text = _get_kb_categories()

    timeout = httpx.Timeout(FS_HTTP_TIMEOUT)

    async def _run():
        async with httpx.AsyncClient(timeout=timeout) as client:
            model_sem = asyncio.Semaphore(FS_MODEL_CONCURRENCY)
            kb_sem = asyncio.Semaphore(FS_KB_CONCURRENCY)

            # ---------- Stage0 场景筛查 ----------
            stage0_result = {
                "relevant": True,
                "confidence": 1.0,
                "scene_tags": [],
                "reason": "disabled",
            }
            if FS_STAGE0_ENABLED:
                stage0_result = await run_stage0_guard_async(
                    client, model_sem, scene_url
                )

            print(f"[Stage0] result = {stage0_result}")

            # 若判定为非建筑/消防场景或置信度不足，直接返回空结果
            if FS_STAGE0_ENABLED and (
                not stage0_result.get("relevant", False)
                or float(stage0_result.get("confidence", 0.0)) < FS_STAGE0_THRESHOLD
            ):
                print("[Stage0] 非建筑/消防相关或置信度不足，短路返回空结果。")
                return {
                    "image_key": image_key,
                    "image_url": scene_url,
                    "stage0": stage0_result,
                    "hazard_types": [],
                    "kb_results": [],
                    "final_result": {"hazard_count": 0, "hazards": []},
                    "type_votes": [],
                    "fusion_votes": [],
                }

            # ---------- StageA 多轮并发猜类型 ----------
            selected_types, type_votes, per_round = await ensemble_stageA_async(
                scene_url,
                allowed_text,
                mapping,
                rounds=FS_STAGEA_ROUNDS,
                temps=FS_TEMP_LIST,
                model_sem=model_sem,
                client=client,
            )

            print("[StageA] 进入后续处理的候选：")
            for i, t in enumerate(selected_types, 1):
                print(f"  {i}. {t['category1']} - {t['category2']}")

            # ---------- KB 并发检索 ----------
            kb_results = await search_kb_for_types_async(
                selected_types, topk=1, kb_sem=kb_sem, client=client
            )

            print("[StageB] KB 命中摘要：")
            for item in kb_results:
                c1, c2 = item["category1"], item["category2"]
                if item.get("results"):
                    hit = (
                        item["results"][0]
                        .get("content", "")[:120]
                        .replace("\n", " ")
                        + "..."
                    )
                else:
                    hit = "(无结果)"
                print(f"  - {c1} - {c2} => {hit}")

            # ---------- Fusion 并发 ----------
            fusion_tasks = []
            for t in selected_types:
                c1, c2 = t["category1"], t["category2"]
                kb_pack = next(
                    (
                        x
                        for x in kb_results
                        if x["category1"] == c1 and x["category2"] == c2
                    ),
                    None,
                )
                if not kb_pack or not kb_pack.get("results"):
                    print(f"[Fusion] 跳过（无KB证据）：{c1} - {c2}")

                    async def _noop(c1=c1, c2=c2):
                        return {
                            "category1": c1,
                            "category2": c2,
                            "present_count": 0,
                            "rounds": FS_FUSION_ROUNDS,
                            "present": False,
                            "hazard": None,
                        }

                    fusion_tasks.append(_noop())
                    continue

                kb_text = kb_pack["results"][0].get("content", "") or ""
                kb_img = (kb_pack["results"][0].get("images") or [None])[0]

                fusion_tasks.append(
                    ensemble_fusion_for_type_async(
                        scene_url,
                        kb_text,
                        kb_img,
                        c1,
                        c2,
                        rounds=FS_FUSION_ROUNDS,
                        model_sem=model_sem,
                        client=client,
                    )
                )

            fusion_res = await asyncio.gather(*fusion_tasks, return_exceptions=True)

        # ---------- 汇总 Fusion ----------
        final_hazards = []
        fusion_votes = []
        for r in fusion_res:
            if isinstance(r, Exception):
                print(f"[Fusion] 任务异常：{r}")
                continue

            fusion_votes.append(
                {
                    "category1": r["category1"],
                    "category2": r["category2"],
                    "present_count": r["present_count"],
                    "rounds": r["rounds"],
                    "present": r["present"],
                }
            )

            status = "存在(多数票)" if r["present"] else "不存在(未过票)"
            print(
                f"[StageC] {r['category1']} - {r['category2']} => {status} {r['present_count']}/{r['rounds']}"
            )

            if r["present"] and r["hazard"]:
                final_hazards.append(r["hazard"])

        final_result = {"hazard_count": len(final_hazards), "hazards": final_hazards}

        return {
            "image_key": image_key,
            "image_url": scene_url,
            "stage0": stage0_result,
            "hazard_types": selected_types,
            "kb_results": kb_results,
            "final_result": final_result,
            "type_votes": type_votes,
            "fusion_votes": fusion_votes,
        }

    # 在后台线程中跑异步流水线
    return asyncio.run(_run())


# ------------------ 后台任务 ------------------
def _process_job(job_id: str, image_key: str):
    """
    后台异步分析任务：
    - 使用 app.py 中的新流水线 Stage0 + StageA + KB + Fusion
    - 正常时将 final_result 写入 result 字段
    - 任意异常（超时 / API 调用错误 / 其它异常）：status=error，result=null
    """
    print(f"[JOB] 开始处理: {job_id}")
    try:
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        # 初始化 / 标记 processing（幂等），顺便清空上一次的 result / error
        collection.update_one(
            {"job_id": job_id},
            {
                "$set": {
                    "status": "processing",
                    "image_key": image_key,
                    "result": None,  # 提交处理时先置空
                    "error": None,
                },
                "$setOnInsert": {"create_time": now_str},
            },
            upsert=True,
        )

        # 调用新流水线
        pipeline_res = run_new_pipeline_sync(image_key)
        final_result = pipeline_res.get(
            "final_result", {"hazard_count": 0, "hazards": []}
        )

        update_doc = {
            "status": "done",
            "result": final_result,  # 成功时写入正常结果
            "finish_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "error": None,
            # 同时把流水线的中间信息也存起来（可选）
            "stage0": pipeline_res.get("stage0"),
            "hazard_types": pipeline_res.get("hazard_types"),
            "kb_results": pipeline_res.get("kb_results"),
            "type_votes": pipeline_res.get("type_votes"),
            "fusion_votes": pipeline_res.get("fusion_votes"),
        }

        collection.update_one(
            {"job_id": job_id},
            {"$set": update_doc},
            upsert=True,
        )

        print(
            f"[JOB] 完成: {job_id} (hazard_count={final_result.get('hazard_count', 0)})"
        )

    except Exception as e:
        # 任意异常（包括超时、API 调用错误等）都走这里
        print(f"[JOB ERROR] {job_id}: {e}")
        collection.update_one(
            {"job_id": job_id},
            {
                "$set": {
                    "status": "error",
                    "result": None,  # 关键：异常时 result 置为 null
                    "error": str(e),
                    "finish_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
            },
            upsert=True,
        )


# ================== 提交与查询接口 ==================
@app.post("/analyze/fire-safety", dependencies=[Depends(verify_token)])
async def submit_analysis(
    background_tasks: BackgroundTasks,
    image: UploadFile = File(...),
    job_id: Optional[str] = Form(None),
):
    """
    提交图片并立即返回:
    - 仅存 key；不存 URL
    - 返回短效预览 URL（前端可立即显示）
    - 后台开始异步分析（使用新流水线）
    """
    try:
        # 生成/接收 job_id
        job_id = job_id or uuid.uuid4().hex
        print(f"[DEBUG] 接收 job_id: {job_id}")

        # 上传到 TOS，得到 key
        image_key = upload_image_to_tos(image)

        # 初始化文档（processing），并清空 result
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        collection.update_one(
            {"job_id": job_id},
            {
                "$set": {
                    "status": "processing",
                    "image_key": image_key,
                    "result": None,  # 提交时就置空 result
                    "error": None,
                },
                "$setOnInsert": {"create_time": now},
            },
            upsert=True,
        )

        # 启动后台任务
        background_tasks.add_task(_process_job, job_id, image_key)

        # 返回短效预览链接（仅用于立刻显示），以及 key（历史由查询接口每次动态签名）
        preview_url = _presigned_get_url(image_key, expires=600)
        return JSONResponse(
            {
                "job_id": job_id,
                "image_key": image_key,
                "image_url": preview_url,
                "status": "processing",
            }
        )

    except HTTPException as he:
        raise he
    except Exception as e:
        print(f"[ERROR] 提交失败: {e}")
        raise HTTPException(status_code=500, detail="提交失败")


@app.get("/analyze/fire-safety/{job_id}", dependencies=[Depends(verify_token)])
async def get_result(job_id: str):
    """
    查询结果（用于前端轮询）
    - 每次查询自动为 image_key 生成一个新鲜的签名URL（避免过期）
    返回:
      - status: processing | done | error | not_found
      - image_key: 仅存 key
      - image_url: 动态签名的短效URL（预览用）
      - result: 正常为 {hazard_count, hazards:[...]}, 异常或处理中为 null
      - (error) error: 错误信息
    """
    url_expires = 600
    doc = collection.find_one({"job_id": job_id})
    if not doc:
        return {"status": "not_found"}

    image_key = doc.get("image_key")
    image_url = _presigned_get_url(image_key, expires=url_expires) if image_key else None

    payload = {
        "status": doc.get("status", "processing"),
        "image_key": image_key,
        "image_url": image_url,  # 每次查询都动态生成，保证不过期
        "create_time": doc.get("create_time"),
        "finish_time": doc.get("finish_time"),
        "result": doc.get("result", None),  # 始终返回 result（可能为 null）
    }
    if doc.get("status") == "error":
        payload["error"] = doc.get("error")

    return JSONResponse(payload)


# ================== 查询历史记录（支持category查询、时间查询、分页） ==================
@app.get("/analyze/history/query", dependencies=[Depends(verify_token)])
async def query_analysis_history(
    category2: Optional[List[str]] = Query(
        None, description="隐患专业二级分类（如 电气专业-应急照明与疏散指示）"
    ),
    start_date: Optional[str] = Query(
        None, description="起始日期，YYYY-MM-DD"
    ),
    end_date: Optional[str] = Query(
        None, description="结束日期，YYYY-MM-DD"
    ),
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
):
    """
    历史记录查询（支持分页、分类、时间段）
    仅统计 status=done 的记录，异常 / 处理中不计入。
    """
    try:
        query = {"status": "done"}

        # 时间筛选
        if start_date or end_date:
            query["create_time"] = {}
            if start_date:
                query["create_time"]["$gte"] = f"{start_date} 00:00:00"
            if end_date:
                query["create_time"]["$lte"] = f"{end_date} 23:59:59"

        # 专业分类筛选（根据 hazards[].category2 含有指定前缀）
        if category2:
            query["result.hazards.category2"] = {
                "$in": [re.compile(f"^{c}") for c in category2]
            }

        total = collection.count_documents(query)
        skip = (page - 1) * limit

        cursor = (
            collection.find(query)
            .sort("create_time", -1)
            .skip(skip)
            .limit(limit)
        )

        results = []
        for doc in cursor:
            preview_url = _presigned_get_url(doc["image_key"], expires=900)
            hazards = doc.get("result", {}).get("hazards", []) or []
            category_set = set()
            for hz in hazards:
                c2 = hz.get("category2", "")
                if c2:
                    category_set.add(c2)

            results.append(
                {
                    "id": str(doc.get("_id")),
                    "job_id": doc.get("job_id", ""),
                    "image_url": preview_url,
                    "hazard_count": doc.get("result", {}).get("hazard_count", 0),
                    "categories": list(category_set),
                    "create_time": doc.get("finish_time") or doc.get("create_time"),
                }
            )

        return {
            "total": total,
            "page": page,
            "limit": limit,
            "data": results,
        }

    except Exception as e:
        print(f"[ERROR] 查询历史记录失败: {e}")
        raise HTTPException(status_code=500, detail="查询失败")


# ================== 批量删除接口 ==================
@app.post("/analyze/fire-safety/delete", dependencies=[Depends(verify_token)])
async def delete_analysis_records(payload: dict = Body(...)):
    """
    批量删除分析记录
    - 接收 JSON: { "job_ids": ["id1", "id2", ...] }
    - 返回: { "deleted_count": n }
    """
    job_ids = payload.get("job_ids", [])
    if not job_ids or not isinstance(job_ids, list):
        raise HTTPException(status_code=400, detail="缺少或格式错误的 job_ids 参数")

    result = collection.delete_many({"job_id": {"$in": job_ids}})
    return {"deleted_count": result.deleted_count}


# ================== 批量下载接口 ==================
@app.post("/analyze/fire-safety/download", dependencies=[Depends(verify_token)])
async def download_reports(payload: dict = Body(...)):
    """
    批量下载报告（Excel，合并同一记录的图片单元格）
    - 入参 JSON: { "job_ids": ["id1", "id2", ...] }；不传则导出全部
    仅导出 status=done 的记录。
    """
    try:
        job_ids = payload.get("job_ids", [])
        query: dict = {"status": "done"}
        if job_ids:
            query["job_id"] = {"$in": job_ids}

        cursor = collection.find(query).sort("create_time", DESCENDING)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "FireSafetyReports"

        # ===== 表头 =====
        ws.append(["序号", "验收图片", "隐患说明", "专业", "违反条例", "分析时间"])

        ws.column_dimensions["A"].width = 6   # 序号
        ws.column_dimensions["B"].width = 22  # 图片列
        ws.column_dimensions["C"].width = 60  # 隐患说明
        ws.column_dimensions["D"].width = 14  # 专业
        ws.column_dimensions["E"].width = 100 # 违反条例
        ws.column_dimensions["F"].width = 20  # 时间

        ws.row_dimensions[1].height = 28
        for col in range(1, 7):
            cell = ws.cell(1, col)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.font = Font(bold=True)

        # ===== 写数据 =====
        row = 2
        seq = 1
        for doc in cursor:
            if not doc or doc.get("status") != "done":
                continue

            hazards = doc.get("result", {}).get("hazards", []) or []

            # 获取图片数据
            img_data = None
            if doc.get("image_key"):
                try:
                    image_url = _presigned_get_url(doc["image_key"], expires=600)
                    img_data = requests.get(image_url, timeout=5).content
                except Exception as e:
                    print(f"[ERROR] 下载图片失败: {e}")

            # === 没有隐患也输出一行 ===
            if not hazards:
                ws.cell(row=row, column=1, value=seq)
                if img_data:
                    try:
                        ximg = XLImage(BytesIO(img_data))
                        ximg.width, ximg.height = 80, 80
                        ws.add_image(ximg, f"B{row}")
                    except Exception as e:
                        print(f"[ERROR] 插入图片失败: {e}")
                        ws.cell(row=row, column=2, value="图片加载失败")
                else:
                    ws.cell(row=row, column=2, value="图片加载失败")

                ws.cell(row=row, column=3, value="未识别到隐患")
                ws.cell(row=row, column=6, value=doc.get("finish_time", ""))
                ws.row_dimensions[row].height = 65
                for c in range(3, 7):
                    ws.cell(row, c).alignment = Alignment(
                        wrap_text=True, vertical="center"
                    )
                row += 1
                seq += 1
                continue

            # === 有隐患：合并 A/B 列 ===
            cnt = len(hazards)
            start_row = row
            end_row = row + cnt - 1
            if cnt > 1:
                ws.merge_cells(
                    start_row=start_row,
                    start_column=1,
                    end_row=end_row,
                    end_column=1,
                )
                ws.merge_cells(
                    start_row=start_row,
                    start_column=2,
                    end_row=end_row,
                    end_column=2,
                )

            # 序号
            ws.cell(start_row, 1, seq).alignment = Alignment(vertical="center")

            # 图片
            if img_data:
                try:
                    ximg = XLImage(BytesIO(img_data))
                    ximg.width, ximg.height = 80, 80
                    ws.add_image(ximg, f"B{start_row}")
                except Exception as e:
                    print(f"[ERROR] 插入图片失败: {e}")
                    ws.cell(start_row, 2, "图片加载失败").alignment = Alignment(
                        vertical="center"
                    )
            else:
                ws.cell(start_row, 2, "图片加载失败").alignment = Alignment(
                    vertical="center"
                )

            # 写 hazards
            for i, hz in enumerate(hazards):
                r = start_row + i
                ws.row_dimensions[r].height = 65

                ws.cell(r, 3, hz.get("description", ""))

                cat1 = hz.get("category1", "")
                # 直接写一级分类作为“专业”列
                major = cat1
                ws.cell(r, 4, major)

                vio_texts = []
                for v in hz.get("violations", []):
                    std = v.get("standard", "")
                    content = v.get("content", "")
                    vio_texts.append(f"{std} {content}")
                ws.cell(r, 5, "\n".join(vio_texts))

                ws.cell(r, 6, doc.get("finish_time", ""))

            # 设置对齐、换行
            for c in range(3, 7):
                for rr in range(start_row, end_row + 1):
                    ws.cell(rr, c).alignment = Alignment(
                        wrap_text=True, vertical="center"
                    )

            row = end_row + 1
            seq += 1

        # ===== 输出 =====
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)

        filename = (
            f"fire_safety_reports_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        )
        headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
        return StreamingResponse(
            stream,
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            headers=headers,
        )

    except HTTPException:
        raise
    except Exception as e:
        print(f"[ERROR] 导出失败: {e}")
        raise HTTPException(status_code=500, detail="导出失败")
